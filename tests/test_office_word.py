from __future__ import annotations

import hashlib
import json
import socket
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from alysis_code.office import word


def digest(value):
    return hashlib.sha256(value).hexdigest()


def make_docx(tmp_path, texts=("Original text",)):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    path = tmp_path / "source.docx"
    doc.save(path)
    return path


def package(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def mutate_package(path, replacements):
    stream = BytesIO()
    with ZipFile(path) as original, ZipFile(stream, "w") as updated:
        updated.comment = b"Preserve this archive comment"
        for item in original.infolist():
            updated.writestr(item, replacements.pop(item.filename, original.read(item)))
        for name, data in replacements.items():
            updated.writestr(name, data)
    path.write_bytes(stream.getvalue())


def selections(source, *edits):
    inspected = word.inspect(source)
    return {
        "expected_source_sha256": inspected["source_sha256"],
        "edits": [
            {
                "locator": f"paragraph:{index}",
                "expected_text_sha256": inspected["paragraphs"][index - 1]["text_sha256"],
                "find": find,
                "replace": replace,
            }
            for index, find, replace in edits
        ],
    }


def test_create_inspect_greek_headings_and_table_without_network(tmp_path, monkeypatch):
    def no_network(*args, **kwargs):
        raise AssertionError("Word helper must not connect to the network")

    monkeypatch.setattr(socket.socket, "connect", no_network)
    output = tmp_path / "report.docx"
    spec = {
        "title": "Εβδομαδιαία έκθεση",
        "blocks": [
            {"kind": "heading", "text": "Σύνοψη", "level": 2},
            {"kind": "paragraph", "text": "Δοκιμή με τόνο και €10."},
            {"kind": "table", "rows": [["Περιγραφή", "Ποσό"], ["Αμοιβή", "10"]]},
        ],
    }
    result = word.create(spec, output, tmp_path)
    reopened = word.inspect(output)
    doc = Document(output)
    assert doc.paragraphs[0].text == spec["title"]
    assert doc.paragraphs[1].style.name == "Heading 2"
    assert doc.tables[0].cell(1, 0).text == "Αμοιβή"
    assert (
        doc.tables[0].rows[0]._tr.find("w:trPr/w:tblHeader", doc.tables[0]._tbl.nsmap) is not None
    )
    assert result["output_sha256"] == reopened["source_sha256"] == digest(output.read_bytes())
    assert reopened["paragraph_count"] == 7
    assert reopened["table_count"] == 1
    assert reopened["coverage"]["status"] == "complete_for_supported_content"
    assert reopened["visual_review"] == "not_performed"
    assert all(item["editable"] for item in reopened["paragraphs"])
    assert reopened["paragraphs"][5]["locator"] == "paragraph:6"
    assert reopened["paragraphs"][5]["text"] == "Αμοιβή"
    json.dumps(result, ensure_ascii=False, allow_nan=False)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["report.docx"]


@pytest.mark.parametrize(
    "spec,message",
    [
        ({"blocks": []}, "1 and 200"),
        ({"blocks": [{"kind": "chart"}]}, "Unsupported block"),
        ({"blocks": [{"kind": "heading", "text": "X", "level": True}]}, "Heading level"),
        ({"blocks": [{"kind": "heading", "text": "X", "level": 10}]}, "Heading level"),
        ({"blocks": [{"kind": "paragraph", "text": "bad\x00"}]}, "unsupported by document XML"),
        ({"blocks": [{"kind": "paragraph", "text": "bad\ud800"}]}, "unsupported by document XML"),
        ({"blocks": [{"kind": "paragraph", "text": "X", "style": "custom"}]}, "Unsupported block"),
        ({"blocks": [{"kind": "table", "rows": [["A", "B"], ["C"]]}]}, "same column"),
        ({"blocks": [{"kind": "table", "rows": [[1]]}]}, "cell must be text"),
        ({"blocks": [{"kind": "table", "rows": [["A"] * 9]}]}, "1-8 columns"),
        (
            {"template": "untrusted.docx", "blocks": [{"kind": "paragraph", "text": "X"}]},
            "Unsupported specification",
        ),
    ],
)
def test_invalid_create_writes_nothing(tmp_path, spec, message):
    with pytest.raises(ValueError, match=message):
        word.create(spec, tmp_path / "invalid.docx", tmp_path)
    assert not list(tmp_path.iterdir())


def test_create_rejects_total_cell_limit_before_allocating_document(tmp_path, monkeypatch):
    def should_not_construct():
        raise AssertionError("Input limits must be validated before constructing a document")

    monkeypatch.setattr(word, "Document", should_not_construct)
    spec = {"blocks": [{"kind": "table", "rows": [[""] * 8] * 200}] * 7}
    with pytest.raises(ValueError, match="10000 main-body paragraphs"):
        word.create(spec, tmp_path / "out.docx", tmp_path)
    assert not list(tmp_path.iterdir())


def test_edit_greek_across_runs_preserves_other_parts_and_paragraphs(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Αγαπητή ")
    paragraph.add_run("Μα").bold = True
    paragraph.add_run("ρία").italic = True
    paragraph.add_run(", καλημέρα.")
    document.add_paragraph("Keep this paragraph")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table source"
    document.sections[0].header.paragraphs[0].text = "Header unchanged"
    document.sections[0].footer.paragraphs[0].text = "Footer unchanged"
    source = tmp_path / "source.docx"
    document.save(source)
    mutate_package(
        source, {"customXml/item99.xml": b'<private xmlns="urn:test">keep exactly</private>'}
    )
    original = source.read_bytes()
    before = package(source)
    before_doc = etree.fromstring(before[word.DOC_PART])
    output = tmp_path / "revision.docx"
    result = word.edit(
        source,
        selections(source, (1, "Μαρία", "Ελένη"), (3, "source", "revised")),
        output,
        tmp_path,
    )
    after = package(output)
    assert source.read_bytes() == original
    assert after.keys() == before.keys()
    assert all(after[name] == value for name, value in before.items() if name != word.DOC_PART)
    after_doc = etree.fromstring(after[word.DOC_PART])
    assert etree.tostring(before_doc.xpath("//w:p", namespaces=word.NS)[1]) == etree.tostring(
        after_doc.xpath("//w:p", namespaces=word.NS)[1]
    )
    revised = Document(output)
    assert revised.paragraphs[0].text == "Αγαπητή Ελένη, καλημέρα."
    assert revised.paragraphs[0].runs[1].text == "Ελένη"
    assert revised.paragraphs[0].runs[1].bold
    assert revised.paragraphs[0].runs[2].text == ""
    assert revised.paragraphs[0].runs[2].italic
    assert revised.tables[0].cell(0, 0).text == "Table revised"
    with ZipFile(output) as archive:
        assert archive.comment == b"Preserve this archive comment"
    assert result["preservation"]["changed_paragraphs"] == ["paragraph:1", "paragraph:3"]
    assert result["preservation"]["package_parts_preserved"]
    assert result["input_sha256"] == digest(original)


@pytest.mark.parametrize(
    "runs,find,replacement,expected",
    [
        (["abc", "def"], "cd", "XY", "abXYef"),
        (["abc", "", "def"], "def", "XYZ", "abcXYZ"),
        (["abc", "def"], "abcdef", "", ""),
        (["a", "b", "c"], "abc", " x & < y ", " x & < y "),
        (["a", "b", "c"], "a", "XYZ", "XYZbc"),
        (["a", "b", "c"], "c", "XYZ", "abXYZ"),
    ],
)
def test_inline_replacement_boundaries(tmp_path, runs, find, replacement, expected):
    document = Document()
    paragraph = document.add_paragraph()
    for value in runs:
        paragraph.add_run(value)
    source = tmp_path / "source.docx"
    document.save(source)
    output = tmp_path / "out.docx"
    word.edit(source, selections(source, (1, find, replacement)), output, tmp_path)
    assert Document(output).paragraphs[0].text == expected


@pytest.mark.parametrize("kind", ["source", "paragraph", "missing_source"])
def test_stale_and_missing_hashes_fail_without_touching_original(tmp_path, kind):
    source = make_docx(tmp_path)
    spec = selections(source, (1, "Original", "New"))
    if kind == "source":
        mutate_package(source, {"customXml/new.xml": b"<new/>"})
    elif kind == "paragraph":
        spec["edits"][0]["expected_text_sha256"] = "0" * 64
    else:
        del spec["expected_source_sha256"]
    before = source.read_bytes()
    output = tmp_path / "out.docx"
    with pytest.raises(ValueError, match="Stale"):
        word.edit(source, spec, output, tmp_path)
    assert source.read_bytes() == before
    assert not output.exists()


@pytest.mark.parametrize(
    "kind,message",
    [
        ("duplicate", "one edit per paragraph"),
        ("missing", "match exactly once"),
        ("ambiguous", "match exactly once"),
        ("newline", "inline text"),
        ("bad_locator", "one-based"),
        ("out_of_bounds", "unavailable"),
        ("unknown_field", "Unsupported edit"),
    ],
)
def test_invalid_edits_publish_no_output(tmp_path, kind, message):
    source = make_docx(tmp_path, ("Original Original", "Another paragraph"))
    spec = selections(source, (1, "Original Original", "Revised"))
    if kind == "duplicate":
        spec["edits"] *= 2
    elif kind in {"missing", "ambiguous"}:
        spec["edits"][0]["find"] = "Absent" if kind == "missing" else "Original"
    elif kind == "newline":
        spec["edits"][0]["replace"] = "New\nparagraph"
    elif kind == "bad_locator":
        spec["edits"][0]["locator"] = "paragraph:01"
    elif kind == "out_of_bounds":
        spec["edits"][0]["locator"] = "paragraph:3"
    else:
        spec["edits"][0]["mode"] = "replace_all"
    before = source.read_bytes()
    output = tmp_path / "out.docx"
    with pytest.raises(ValueError, match=message):
        word.edit(source, spec, output, tmp_path)
    assert not output.exists()
    assert source.read_bytes() == before


@pytest.mark.parametrize(
    "complex_xml",
    [
        '<w:hyperlink w:anchor="target"><w:r><w:t>Complex</w:t></w:r></w:hyperlink>',
        '<w:r><w:fldChar w:fldCharType="begin"/><w:t>Complex</w:t></w:r>',
        "<w:r><w:t>Complex</w:t><w:drawing/></w:r>",
        '<w:bookmarkStart w:id="0" w:name="keep"/><w:r><w:t>Complex</w:t></w:r><w:bookmarkEnd w:id="0"/>',
        "<w:r><w:t>Complex</w:t><w:tab/></w:r>",
        '<w:r><w:t>Complex</w:t></w:r><m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x</m:t></m:r></m:oMath>',
    ],
)
def test_complex_selection_rejected_but_other_paragraph_edit_preserves_it(tmp_path, complex_xml):
    source = make_docx(tmp_path, ("Simple", "Complex"))
    parts = package(source)
    root = etree.fromstring(parts[word.DOC_PART])
    paragraphs = root.xpath("//w:p", namespaces=word.NS)
    complex_paragraph = etree.fromstring(f'<w:p xmlns:w="{word.W}">{complex_xml}</w:p>')
    paragraphs[1].getparent().replace(paragraphs[1], complex_paragraph)
    mutate_package(source, {word.DOC_PART: etree.tostring(root)})
    inspected = word.inspect(source)
    assert inspected["paragraphs"][0]["editable"]
    assert not inspected["paragraphs"][1]["editable"]
    assert inspected["coverage"]["status"] == "partial"
    with pytest.raises(ValueError, match="complex content|non-text run"):
        word.edit(
            source, selections(source, (2, "Complex", "New")), tmp_path / "bad.docx", tmp_path
        )
    output = tmp_path / "good.docx"
    word.edit(source, selections(source, (1, "Simple", "Revised")), output, tmp_path)
    assert word.inspect(output)["paragraphs"][1] == inspected["paragraphs"][1]
    assert not (tmp_path / "bad.docx").exists()


def test_content_control_ancestor_is_not_editable(tmp_path):
    source = make_docx(tmp_path)
    parts = package(source)
    root = etree.fromstring(parts[word.DOC_PART])
    paragraph = root.find("w:body/w:p", word.NS)
    body = paragraph.getparent()
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    body.replace(paragraph, control)
    content.append(paragraph)
    control.append(content)
    mutate_package(source, {word.DOC_PART: etree.tostring(root)})
    assert not word.inspect(source)["paragraphs"][0]["editable"]
    with pytest.raises(ValueError, match="inside unsupported complex"):
        word.edit(
            source, selections(source, (1, "Original", "Revised")), tmp_path / "out.docx", tmp_path
        )


@pytest.mark.parametrize("restriction", ["tracked", "protected"])
def test_document_restrictions_are_visible_and_prevent_edits(tmp_path, restriction):
    source = make_docx(tmp_path)
    parts = package(source)
    if restriction == "tracked":
        root = etree.fromstring(parts[word.DOC_PART])
        root.find("w:body", word.NS).append(OxmlElement("w:ins"))
        mutate_package(source, {word.DOC_PART: etree.tostring(root)})
    else:
        root = etree.fromstring(parts["word/settings.xml"])
        protection = OxmlElement("w:documentProtection")
        protection.set(qn("w:enforcement"), "true")
        root.append(protection)
        mutate_package(source, {"word/settings.xml": etree.tostring(root)})
    result = word.inspect(source)
    assert not result["paragraphs"][0]["editable"]
    assert result["coverage"]["status"] == "partial"
    with pytest.raises(ValueError, match="Tracked changes|protection"):
        word.edit(
            source, selections(source, (1, "Original", "New")), tmp_path / "out.docx", tmp_path
        )


@pytest.mark.parametrize(
    "part",
    [
        "word/vbaProject.bin",
        "word/embeddings/file.bin",
        "word/activeX/1.xml",
        "_xmlsignatures/sig1.xml",
    ],
)
def test_unsupported_active_packages_rejected(tmp_path, part):
    source = make_docx(tmp_path)
    mutate_package(source, {part: b"unsupported"})
    with pytest.raises(ValueError, match="unsupported"):
        word.inspect(source)


def test_entity_declaration_in_utf16_xml_is_rejected(tmp_path):
    source = make_docx(tmp_path)
    xml = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE a [<!ENTITY example "data">]><a>&example;</a>'
    mutate_package(source, {"customXml/entity.xml": xml.encode("utf-16")})
    with pytest.raises(ValueError, match="document type declarations"):
        word.inspect(source)


def test_output_never_overwrites_source_or_existing_file(tmp_path):
    source = make_docx(tmp_path)
    before = source.read_bytes()
    with pytest.raises(FileExistsError):
        word.edit(source, selections(source, (1, "Original", "New")), source, tmp_path)
    with pytest.raises(FileExistsError):
        word.create({"blocks": [{"kind": "paragraph", "text": "New"}]}, source, tmp_path)
    assert source.read_bytes() == before


def test_inspection_preserves_tabs_breaks_and_does_not_repeat_nested_paragraph_text(tmp_path):
    source = make_docx(tmp_path, ("A\tB\nC",))
    root = etree.fromstring(package(source)[word.DOC_PART])
    outer = root.find("w:body/w:p", word.NS)
    drawing = etree.fromstring(
        f'<w:r xmlns:w="{word.W}"><w:drawing><w:txbxContent><w:p><w:r><w:t>Nested</w:t></w:r></w:p></w:txbxContent></w:drawing></w:r>'
    )
    outer.append(drawing)
    mutate_package(source, {word.DOC_PART: etree.tostring(root)})
    result = word.inspect(source)
    assert [p["text"] for p in result["paragraphs"]] == ["A\tB\nC", "Nested"]
    assert not any(p["editable"] for p in result["paragraphs"])
    assert result["coverage"]["status"] == "partial"
