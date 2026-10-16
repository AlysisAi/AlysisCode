from __future__ import annotations

import hashlib
import io
import json
import os
import signal
import subprocess
import sys
import time
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from alysis_code.office import office_capability_status
from alysis_code.office.__main__ import _package, _spec, run


def args(root, command="create", **values):
    return SimpleNamespace(
        workspace=str(root),
        command=command,
        format="docx",
        spec="request.json",
        source=None,
        output="result.docx",
        **values,
    )


def test_real_command_creates_editable_file_and_rejects_overwrite(tmp_path):
    request = {
        "blocks": [
            {"kind": "heading", "text": "Ελληνική αναφορά"},
            {"kind": "paragraph", "text": "Original content."},
        ]
    }
    (tmp_path / "request.json").write_text(json.dumps(request), encoding="utf-8")
    created = run(args(tmp_path))
    assert created["status"] == "created"
    assert Document(tmp_path / "result.docx").paragraphs[0].text == "Ελληνική αναφορά"
    original = (tmp_path / "result.docx").read_bytes()
    assert created["sha256"] == hashlib.sha256(original).hexdigest()
    assert created["preview"]["visual_review"] == "not_performed"
    with pytest.raises(ValueError, match="new output"):
        run(args(tmp_path))
    assert (tmp_path / "result.docx").read_bytes() == original
    assert not list(tmp_path.glob(".alysis-office-*"))


def test_command_rejects_symlink_and_escape_outputs_before_writing(tmp_path):
    (tmp_path / "request.json").write_text('{"blocks":[{"kind":"paragraph","text":"sample"}]}')
    link = tmp_path / "linked"
    link.symlink_to(tmp_path, target_is_directory=True)
    for output in ("linked/result.docx", "../escape.docx"):
        request = args(tmp_path)
        request.output = output
        with pytest.raises(ValueError, match="symbolic|inside"):
            run(request)


def test_missing_renderer_does_not_modify_editable_file(tmp_path, monkeypatch):
    monkeypatch.delenv("ALYSIS_OFFICE_SOFFICE", raising=False)
    source = tmp_path / "original.docx"
    Document().save(source)
    original = source.read_bytes()
    with pytest.raises(ValueError, match="preview is unavailable"):
        run(
            SimpleNamespace(
                workspace=str(tmp_path),
                command="render",
                source="original.docx",
                preview_dir="preview",
            )
        )
    assert source.read_bytes() == original
    assert not (tmp_path / "preview").exists()
    assert not list(tmp_path.glob(".alysis-office-*"))


def test_stale_edit_fails_before_adapter_and_leaves_original(tmp_path):
    source = tmp_path / "original.docx"
    Document().save(source)
    original = source.read_bytes()
    (tmp_path / "request.json").write_text('{"expected_source_sha256":"stale","edits":[]}')
    with pytest.raises(ValueError, match="Source changed"):
        run(
            SimpleNamespace(
                workspace=str(tmp_path),
                command="edit",
                source="original.docx",
                spec="request.json",
                output="edited.docx",
            )
        )
    assert source.read_bytes() == original
    assert not (tmp_path / "edited.docx").exists()


@pytest.mark.parametrize("name", ["../evil.xml", "word/vbaProject.bin", "/evil.xml"])
def test_invalid_office_package_rejected(name):
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(name, "<data/>")
    with pytest.raises(ValueError):
        _package(data.getvalue())


@pytest.mark.parametrize("text", ['{"value":NaN}', '{"a":1,"a":2}', "[]"])
def test_request_rejects_ambiguous_json(tmp_path, text):
    path = tmp_path / "request.json"
    path.write_text(text)
    with pytest.raises(ValueError):
        _spec(path)


def test_capability_names_exact_interpreter_and_never_claims_rendered(monkeypatch):
    monkeypatch.delenv("ALYSIS_OFFICE_SOFFICE", raising=False)
    status = office_capability_status()
    assert status["state"] == "ready"
    assert sys.executable in status["command"]
    assert status["preview"]["state"] == "unavailable"
    assert "not calculated" in status["limits"]
    assert "permissions" in status["guidance"]


def test_cli_describe_and_json_error(tmp_path):
    environment = dict(os.environ, PYTHONPATH=str(Path(__file__).parents[1] / "src"))
    for format_name in ("docx", "xlsx", "pptx"):
        result = subprocess.run(
            [sys.executable, "-m", "alysis_code.office", "describe", "--format", format_name],
            env=environment,
            capture_output=True,
            text=True,
            cwd=tmp_path,
        )
        assert result.returncode == 0, result.stderr
        assert json.loads(result.stdout)["guide"]
    result = subprocess.run(
        [sys.executable, "-m", "alysis_code.office", "inspect", "--source", "missing.docx"],
        env=environment,
        capture_output=True,
        text=True,
        cwd=tmp_path,
    )
    assert result.returncode == 1
    assert json.loads(result.stdout)["status"] == "error"


@pytest.mark.skipif(sys.platform not in {"darwin", "linux"}, reason="POSIX Office renderer")
def test_preview_directory_publication_is_complete_and_exclusive(tmp_path):
    from alysis_code.office.rendering import _publish_directory

    stage = tmp_path / "stage"
    stage.mkdir()
    (stage / "page.png").write_bytes(b"complete")
    destination = tmp_path / "preview"
    destination.mkdir()
    with pytest.raises(FileExistsError):
        _publish_directory(stage, destination)
    assert stage.exists() and list(destination.iterdir()) == []
    destination.rmdir()
    _publish_directory(stage, destination)
    assert not stage.exists()
    assert (destination / "page.png").read_bytes() == b"complete"


@pytest.mark.skipif(os.name != "posix", reason="POSIX renderer process group")
def test_converter_kills_child_even_when_launcher_exits(tmp_path):
    import psutil

    from alysis_code.office.rendering import _convert

    pid_file = tmp_path / "child.pid"
    launcher = tmp_path / "launcher"
    launcher.write_text(
        "#!" + sys.executable + "\nimport subprocess,sys,pathlib\n"
        "p=subprocess.Popen([sys.executable,'-c','import time; time.sleep(30)'])\n"
        f"pathlib.Path({str(pid_file)!r}).write_text(str(p.pid))\n"
    )
    launcher.chmod(0o700)
    converted = tmp_path / "converted"
    converted.mkdir()
    _convert(launcher, tmp_path / "unused.docx", converted, tmp_path / "profile")
    pid = int(pid_file.read_text())
    try:
        for _ in range(50):
            if not psutil.pid_exists(pid) or psutil.Process(pid).status() == psutil.STATUS_ZOMBIE:
                break
            time.sleep(0.01)
        else:
            pytest.fail("Converter descendant survived its owning process group")
    finally:
        if psutil.pid_exists(pid):
            psutil.Process(pid).kill()


@pytest.mark.skipif(os.name != "posix", reason="POSIX cancellation")
def test_signal_during_preview_publication_leaves_no_partial_result(tmp_path, monkeypatch):
    from reportlab.pdfgen.canvas import Canvas

    from alysis_code.office import rendering

    source = tmp_path / "source.docx"
    Document().save(source)
    original = source.read_bytes()
    monkeypatch.setenv("ALYSIS_OFFICE_SOFFICE", sys.executable)

    def convert(executable, source, destination, profile):
        canvas = Canvas(str(destination / (source.stem + ".pdf")))
        canvas.drawString(60, 700, "Test page")
        canvas.save()

    def interrupted(source, destination):
        os.kill(os.getpid(), signal.SIGTERM)

    monkeypatch.setattr(rendering, "_convert", convert)
    monkeypatch.setattr(rendering, "_publish_directory", interrupted)
    with pytest.raises(KeyboardInterrupt):
        run(
            SimpleNamespace(
                workspace=str(tmp_path),
                command="render",
                source="source.docx",
                preview_dir="preview",
            )
        )
    assert not (tmp_path / "preview").exists()
    assert not list(tmp_path.glob(".alysis-office-*"))
    assert source.read_bytes() == original
