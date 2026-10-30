from __future__ import annotations

import hashlib
import json
import os
import threading
import zipfile
from contextlib import closing
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from types import SimpleNamespace

import openpyxl
import pytest
import xlsxwriter

from alysis_code.office.__main__ import run


def calculate(root, source="source.xlsx", output="calculated.xlsx"):
    return run(
        SimpleNamespace(workspace=str(root), command="recalculate", source=source, output=output)
    )


def test_unavailable_calculator_preserves_original(tmp_path, monkeypatch):
    monkeypatch.setattr("alysis_code.office.calculation.renderer_path", lambda: None)
    with xlsxwriter.Workbook(tmp_path / "source.xlsx") as book:
        book.add_worksheet().write_formula("A1", "=2+2", None, 999)
    original = (tmp_path / "source.xlsx").read_bytes()
    with pytest.raises(ValueError, match="unavailable"):
        calculate(tmp_path)
    assert (tmp_path / "source.xlsx").read_bytes() == original
    assert not (tmp_path / "calculated.xlsx").exists()


def test_authoring_does_not_restrict_function_names_to_a_handwritten_list(tmp_path):
    spec = {"sheets": [{"name": "Data", "rows": [[1], [4], [7], [{"formula": "=MEDIAN(A1:A3)"}]]}]}
    (tmp_path / "spec.json").write_text(json.dumps(spec))
    result = run(
        SimpleNamespace(
            workspace=str(tmp_path),
            command="create",
            format="xlsx",
            spec="spec.json",
            output="source.xlsx",
        )
    )
    assert result["content"]["formula_count"] == 1
    assert result["content"]["formula_results"] == "not_calculated"


@pytest.fixture
def real_calculator(monkeypatch):
    executable = os.environ.get("ALYSIS_TEST_OFFICE_SOFFICE")
    if not executable:
        pytest.skip("Set ALYSIS_TEST_OFFICE_SOFFICE to exercise an installed real calculator")
    monkeypatch.setenv("ALYSIS_OFFICE_SOFFICE", executable)


def test_real_calculation_replaces_stale_results_preserves_package_and_source(
    tmp_path, real_calculator
):
    source = tmp_path / "source.xlsx"
    with xlsxwriter.Workbook(source) as book:
        sheet = book.add_worksheet("Έσοδα")
        for row, value in enumerate((13, 7, 11)):
            sheet.write(row, 0, value)
        sheet.write_formula("A4", "=SUM(A1:A3)", None, 999)
        sheet.write_formula("B1", '=IF(A1>0,"ok","")', None, "stale")
        sheet.write_formula("B2", '=IF(A2>0,"","")', None, "stale")
        sheet.write_formula("B3", "=A4=31", None, False)
        sheet.write_formula("B4", "=MEDIAN(A1:A3)", None, -999)
        book.add_worksheet("Summary").write_formula("A1", "='Έσοδα'!A4*2", None, -999)
    original = source.read_bytes()
    result = calculate(tmp_path)
    assert result["calculation"]["state"] == "passed"
    assert result["calculation"]["formula_count"] == 6
    assert result["content"]["needs_recalculation"] is False
    output = tmp_path / "calculated.xlsx"
    assert result["sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert source.read_bytes() == original
    with closing(openpyxl.load_workbook(output, data_only=True)) as book:
        # openpyxl maps a valid empty-string <v/> cache to None. The operation's
        # typed result still distinguishes it from a missing cache.
        assert [book["Έσοδα"][cell].value for cell in ("A4", "B1", "B2", "B3", "B4")] == [
            31,
            "ok",
            None,
            True,
            11,
        ]
        assert book["Summary"]["A1"].value == 62
    assert next(item for item in result["calculation"]["results"] if item["cell"] == "B2") == {
        "sheet": "Έσοδα",
        "cell": "B2",
        "type": "str",
        "value": "",
    }
    with (
        closing(openpyxl.load_workbook(source)) as before,
        closing(openpyxl.load_workbook(output)) as after,
    ):
        assert {
            sheet.title: [(c.coordinate, c.value, c.style_id) for row in sheet for c in row]
            for sheet in before
        } == {
            sheet.title: [(c.coordinate, c.value, c.style_id) for row in sheet for c in row]
            for sheet in after
        }
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if not name.startswith("xl/worksheets/"):
                assert before.read(name) == after.read(name)


def test_real_calculation_reports_errors(tmp_path, real_calculator):
    with xlsxwriter.Workbook(tmp_path / "source.xlsx") as book:
        book.add_worksheet().write_formula("A1", "=1/0", None, 42)
    result = calculate(tmp_path)
    assert result["calculation"]["state"] == "failed"
    assert result["calculation"]["errors"][0]["value"] == "#DIV/0!"
    assert result["content"]["needs_recalculation"] is True


def test_real_calculation_preserves_layout_and_metadata_that_editing_cannot_roundtrip(
    tmp_path, real_calculator
):
    from alysis_code.artifact_checks import ArtifactChecks
    from alysis_code.office import sheets

    source = tmp_path / "source.xlsx"
    with xlsxwriter.Workbook(source) as book:
        book.set_custom_property("Review status", "Draft – Έλεγχος")
        sheet = book.add_worksheet("Δεδομένα")
        sheet.repeat_rows(0)
        sheet.repeat_columns(0)
        sheet.print_area("A1:C3")
        sheet.set_landscape()
        sheet.write_row("A1", ["Input", "Result", "Note"])
        sheet.write("A2", 21)
        sheet.write_formula("B2", "=A2*2", None, 999)
        sheet.conditional_format(
            "B2",
            {
                "type": "cell",
                "criteria": ">",
                "value": 40,
                "format": book.add_format({"bg_color": "#e4f5df"}),
            },
        )
    original = source.read_bytes()
    assert not sheets.inspect(source)["editable"]
    # LibreOffice writes this ordinary-formula recalculation flag on save.
    from alysis_code.office.calculation import NS, _write_parts, _xml

    with zipfile.ZipFile(source) as archive:
        tree = _xml(archive.read("xl/worksheets/sheet1.xml"))
        tree.find(f".//{{{NS}}}f").set("aca", "false")
        from lxml import etree

        _write_parts(
            original, {"xl/worksheets/sheet1.xml": etree.tostring(tree)}, tmp_path / "flagged.xlsx"
        )
    source.write_bytes((tmp_path / "flagged.xlsx").read_bytes())
    original = source.read_bytes()
    result = calculate(tmp_path)
    assert result["calculation"]["state"] == "passed"
    with closing(openpyxl.load_workbook(tmp_path / "calculated.xlsx", data_only=True)) as book:
        assert book.active["B2"].value == 42
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(tmp_path / "calculated.xlsx") as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert before.read(name) == after.read(name)
    assert source.read_bytes() == original
    report = ArtifactChecks(tmp_path).check("calculated.xlsx")
    assert (
        next(item for item in report["items"] if item["id"] == "calculation")["status"] == "passed"
    )


def test_real_calculation_never_fetches_external_formula_content(tmp_path, real_calculator):
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"external-content")

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with xlsxwriter.Workbook(tmp_path / "source.xlsx") as book:
            book.add_worksheet().write_formula(
                "A1", f'=WEBSERVICE("http://127.0.0.1:{server.server_port}/probe")', None, "stale"
            )
        result = calculate(tmp_path)
        assert requests == []
        assert result["calculation"]["state"] == "failed"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_calculation_rejects_converter_shared_string_indexes(tmp_path, monkeypatch):
    from alysis_code.office import calculation

    source = tmp_path / "source.xlsx"
    with xlsxwriter.Workbook(source) as book:
        sheet = book.add_worksheet()
        sheet.write_string("A1", "original string table")
        sheet.write_formula("B1", "=A1", None, "stale")
    original = source.read_bytes()
    with zipfile.ZipFile(source) as archive:
        tree = calculation._xml(archive.read("xl/worksheets/sheet1.xml"))
    cell = next(c for c in tree.iter(f"{{{calculation.NS}}}c") if c.get("r") == "B1")
    cell.set("t", "s")
    cell.find(f"{{{calculation.NS}}}v").text = "0"

    def convert(_executable, _input, converted, _profile, **_options):
        calculation._write_parts(
            original,
            {"xl/worksheets/sheet1.xml": calculation.ET.tostring(tree)},
            converted / "input.xlsx",
        )

    monkeypatch.setattr(calculation, "renderer_path", lambda: tmp_path / "calculator")
    monkeypatch.setattr(calculation, "_convert", convert)
    with pytest.raises(ValueError, match="unsupported cache type"):
        calculate(tmp_path)
    assert source.read_bytes() == original
    assert not (tmp_path / "calculated.xlsx").exists()
