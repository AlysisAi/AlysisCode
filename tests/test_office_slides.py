from __future__ import annotations

import hashlib
import json
import socket
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from alysis_code.office import slides


def digest(value):
    return hashlib.sha256(value.encode()).hexdigest()


def create_deck(root, spec=None):
    path = root / "source.pptx"
    result = slides.create(
        spec
        or {
            "slides": [
                {"title": "Αρχικός τίτλος", "body": ["Πρώτη παράγραφος", "Δεύτερη παράγραφος"]}
            ]
        },
        path,
        root,
    )
    return path, result


def change_for(result, *, shape=0, text="Νέος τίτλος", **kwargs):
    item = result["slides"][0]["shapes"][shape]
    return {
        "expected_source_sha256": result["source_sha256"],
        "edits": [
            {
                "slide": 1,
                "shape_id": item["shape_id"],
                "expected_text_sha256": item["text_sha256"],
                "text": text,
                **kwargs,
            }
        ],
    }


def test_create_inspect_greek_tables_images_and_no_network(tmp_path, monkeypatch):
    def no_network(*args, **kwargs):
        raise AssertionError("PPTX operations must not access the network")

    monkeypatch.setattr(socket.socket, "connect", no_network)
    Image.new("RGB", (200, 100), "red").save(tmp_path / "picture.png")
    spec = {
        "title": "Ελληνική παρουσίαση",
        "slides": [
            {
                "title": "Προγραμματισμός",
                "elements": [
                    {
                        "type": "text",
                        "text": "Αθήνα και Θεσσαλονίκη",
                        "x": 0.6,
                        "y": 1.5,
                        "width": 6,
                        "height": 1,
                    },
                    {
                        "type": "table",
                        "rows": [["Πόλη", "Ποσό"], ["Αθήνα", "42"]],
                        "x": 0.6,
                        "y": 3,
                        "width": 6,
                        "height": 2,
                    },
                    {
                        "type": "image",
                        "path": "picture.png",
                        "x": 7,
                        "y": 1.5,
                        "width": 5,
                        "height": 4,
                    },
                ],
            }
        ],
    }
    path, result = create_deck(tmp_path, spec)
    assert result["status"] == "created"
    assert result["slide_count"] == 1
    assert result["visual_review"] == "not_performed"
    assert result["text_overflow"] == "not_verified"
    assert result["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    items = result["slides"][0]["shapes"]
    assert not any(item["outside_slide"] for item in items)
    assert items[0]["text"] == "Προγραμματισμός"
    assert items[2]["cells"][2]["text"] == "Αθήνα"
    assert items[3]["text_editable"] is False
    picture = Presentation(path).slides[0].shapes[3]
    assert picture.width == picture.height * 2
    assert Presentation(path).core_properties.title == "Ελληνική παρουσίαση"
    json.dumps(result)


def test_edit_preserves_source_other_slide_chart_media_and_style(tmp_path):
    source, _ = create_deck(tmp_path)
    presentation = Presentation(source)
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Counts", [2, 3])
    presentation.slides[0].shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(3), Inches(4), Inches(3), data
    )
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2)).text = "Unchanged slide"
    Image.new("RGB", (32, 32)).save(tmp_path / "image.png")
    second.shapes.add_picture(str(tmp_path / "image.png"), Inches(6), Inches(1), Inches(1))
    presentation.save(source)
    original = source.read_bytes()
    before = slides.inspect(source)
    output = tmp_path / "revision.pptx"
    result = slides.edit(source, change_for(before), output, tmp_path)
    assert source.read_bytes() == original
    assert result["slides"][0]["shapes"][0]["text"] == "Νέος τίτλος"
    assert result["changed_package_parts"] == ["ppt/slides/slide1.xml"]
    edited = Presentation(output)
    assert edited.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.bold
    assert edited.slides[0].shapes[2].chart.series[0].values == (2.0, 3.0)
    with ZipFile(source) as before_zip, ZipFile(output) as after_zip:
        assert before_zip.namelist() == after_zip.namelist()
        for name in before_zip.namelist():
            if name != "ppt/slides/slide1.xml":
                assert before_zip.read(name) == after_zip.read(name), name


def test_precise_run_edit_keeps_formatting_links_and_surrounding_text(tmp_path):
    source, _ = create_deck(tmp_path)
    presentation = Presentation(source)
    paragraph = presentation.slides[0].shapes[0].text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = " linked"
    run.font.italic = True
    run.hyperlink.address = "https://example.com/"
    presentation.save(source)
    result = slides.inspect(source)
    assert not result["slides"][0]["shapes"][0]["whole_text_editable"]
    spec = change_for(result, text=" revised", paragraph=1, run=2)
    spec["edits"][0]["expected_text_sha256"] = digest(" linked")
    output = tmp_path / "revised.pptx"
    slides.edit(source, spec, output, tmp_path)
    paragraph = Presentation(output).slides[0].shapes[0].text_frame.paragraphs[0]
    assert paragraph.text == "Αρχικός τίτλος revised"
    assert paragraph.runs[0].font.bold
    assert paragraph.runs[1].font.italic
    assert paragraph.runs[1].hyperlink.address == "https://example.com/"


@pytest.mark.parametrize(
    "kind",
    [
        "source",
        "text",
        "complex",
        "unknown_shape",
        "same_output",
        "overlap",
        "paragraph_break",
    ],
)
def test_invalid_edits_preserve_original_without_output(tmp_path, kind):
    source, result = create_deck(tmp_path)
    original = source.read_bytes()
    output = tmp_path / "revision.pptx"
    spec = change_for(result)
    if kind == "source":
        spec["expected_source_sha256"] = "0" * 64
    elif kind == "text":
        spec["edits"][0]["expected_text_sha256"] = "0" * 64
    elif kind == "complex":
        spec = change_for(result, shape=1)
    elif kind == "unknown_shape":
        spec["edits"][0]["shape_id"] = 900
    elif kind == "same_output":
        output = source
    elif kind == "overlap":
        spec["edits"] *= 2
    elif kind == "paragraph_break":
        spec["edits"][0]["text"] = "Two\nparagraphs"
    with pytest.raises(ValueError):
        slides.edit(source, spec, output, tmp_path)
    assert source.read_bytes() == original
    if output != source:
        assert not output.exists()


def test_table_cell_edit_and_merged_cell_rejection(tmp_path):
    source, result = create_deck(
        tmp_path,
        {
            "slides": [
                {
                    "elements": [
                        {
                            "type": "table",
                            "rows": [["Name", "Value"], ["Alice", "42"]],
                            "x": 1,
                            "y": 1,
                            "width": 8,
                            "height": 3,
                        }
                    ]
                }
            ]
        },
    )
    item = result["slides"][0]["shapes"][0]
    spec = {
        "expected_source_sha256": result["source_sha256"],
        "edits": [
            {
                "slide": 1,
                "shape_id": item["shape_id"],
                "row": 2,
                "column": 2,
                "text": "84",
                "expected_text_sha256": digest("42"),
            }
        ],
    }
    output = tmp_path / "revised.pptx"
    slides.edit(source, spec, output, tmp_path)
    assert Presentation(output).slides[0].shapes[0].table.cell(1, 1).text == "84"
    presentation = Presentation(source)
    presentation.slides[0].shapes[0].table.cell(1, 0).merge(
        presentation.slides[0].shapes[0].table.cell(1, 1)
    )
    presentation.save(source)
    spec["expected_source_sha256"] = slides.inspect(source)["source_sha256"]
    with pytest.raises(ValueError, match="Merged"):
        slides.edit(source, spec, tmp_path / "invalid.pptx", tmp_path)


@pytest.mark.parametrize(
    "element,reason",
    [
        ({"type": "text", "text": "Outside", "x": 12, "y": 1, "width": 2, "height": 1}, "outside"),
        (
            {"type": "text", "text": "A", "x": True, "y": 1, "width": 5, "height": 1},
            "finite number",
        ),
        (
            {"type": "text", "text": "A", "x": 1, "y": 1, "width": float("nan"), "height": 1},
            "finite number",
        ),
        ({"type": "text", "text": "A\x00B", "x": 1, "y": 1, "width": 5, "height": 1}, "control"),
        ({"type": "chart"}, "Supported element"),
        (
            {"type": "table", "rows": [["A"], ["B", "C"]], "x": 1, "y": 1, "width": 5, "height": 3},
            "equal column",
        ),
    ],
)
def test_invalid_creation_has_no_output(tmp_path, element, reason):
    output = tmp_path / "output.pptx"
    with pytest.raises(ValueError, match=reason):
        slides.create({"slides": [{"elements": [element]}]}, output, tmp_path)
    assert not output.exists()


@pytest.mark.parametrize(
    "path",
    ["../outside.png", "/tmp/outside.png", "https://example.com/a.png", "link.png", "broken.png"],
)
def test_images_reject_unsafe_or_invalid_paths(tmp_path, path, require_symlinks):
    Image.new("RGB", (2, 2)).save(tmp_path / "safe.png")
    (tmp_path / "link.png").symlink_to(tmp_path / "safe.png")
    (tmp_path / "broken.png").write_bytes(b"not an image")
    with pytest.raises(ValueError):
        slides.create(
            {
                "slides": [
                    {
                        "elements": [
                            {"type": "image", "path": path, "x": 1, "y": 1, "width": 2, "height": 2}
                        ]
                    }
                ]
            },
            tmp_path / "output.pptx",
            tmp_path,
        )
    assert not (tmp_path / "output.pptx").exists()


def test_dynamic_fields_are_not_flattened_and_groups_are_not_editable(tmp_path):
    source, _ = create_deck(tmp_path)
    presentation = Presentation(source)
    paragraph = presentation.slides[0].shapes[0].text_frame.paragraphs[0]
    field = OxmlElement("a:fld")
    field.set("id", "{00000000-0000-0000-0000-000000000001}")
    field.set("type", "slidenum")
    content = OxmlElement("a:t")
    content.text = "1"
    field.append(content)
    paragraph._p.append(field)
    group = presentation.slides[0].shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text = "Group text"
    presentation.save(source)
    result = slides.inspect(source)
    assert not result["slides"][0]["shapes"][0]["whole_text_editable"]
    assert "Grouped" in result["slides"][0]["shapes"][2]["unsupported_edit_reason"]
    with pytest.raises(ValueError, match="Complex"):
        slides.edit(source, change_for(result), tmp_path / "revised.pptx", tmp_path)


def test_inspect_reports_existing_shape_outside_slide(tmp_path):
    source, _ = create_deck(tmp_path)
    presentation = Presentation(source)
    presentation.slides[0].shapes[0].left = Inches(-1)
    presentation.save(source)
    assert slides.inspect(source)["slides"][0]["shapes"][0]["outside_slide"]


def test_slide_limit_unknown_fields_and_invalid_package(tmp_path):
    with pytest.raises(ValueError, match="slides"):
        slides.create({"slides": [{"title": "A"}] * 51}, tmp_path / "x.pptx", tmp_path)
    with pytest.raises(ValueError, match="Unsupported"):
        slides.create(
            {"slides": [{"title": "A"}], "transitions": True}, tmp_path / "x.pptx", tmp_path
        )
    (tmp_path / "bad.pptx").write_bytes(b"not zip")
    with pytest.raises(ValueError, match="Cannot read"):
        slides.inspect(tmp_path / "bad.pptx")


def test_inspection_rejects_excess_empty_text_nodes(tmp_path):
    source, _ = create_deck(tmp_path)
    presentation = Presentation(source)
    paragraph = presentation.slides[0].shapes[0].text_frame.paragraphs[0]
    for _ in range(20_001):
        paragraph.add_run()
    presentation.save(source)
    with pytest.raises(ValueError, match="inspection limits"):
        slides.inspect(source)


def test_metadata_title_is_not_silently_truncated(tmp_path):
    with pytest.raises(ValueError, match="255"):
        slides.create(
            {"title": "A" * 256, "slides": [{"title": "Hello"}]}, tmp_path / "output.pptx", tmp_path
        )
    assert not (tmp_path / "output.pptx").exists()


def test_layout_is_not_decided_from_character_width_estimates(tmp_path):
    text = "iiiiiiiiiiiiiiiiiiii Ω δοκιμή 漢字 " * 80
    source, result = create_deck(
        tmp_path,
        {
            "slides": [
                {
                    "elements": [
                        {"type": "text", "text": text, "x": 1, "y": 1, "width": 5, "height": 1}
                    ]
                }
            ]
        },
    )
    assert Presentation(source).slides[0].shapes[0].text == text
    assert result["text_overflow"] == "not_verified"
    assert result["visual_review"] == "not_performed"
