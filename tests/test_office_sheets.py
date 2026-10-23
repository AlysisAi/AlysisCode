from __future__ import annotations

import hashlib
import json
import socket
import zipfile
from contextlib import closing
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
import xlsxwriter

from alysis_code.office import sheets


def create(root: Path, spec: dict | None = None) -> tuple[Path, dict]:
    source = root / "source.xlsx"
    result = sheets.create(spec or sheets.example_spec(), source, root)
    return source, result


def edit(root: Path, source: Path, digest: str, changes: list) -> tuple[Path, dict]:
    output = root / "revised.xlsx"
    return output, sheets.edit(
        source, {"expected_source_sha256": digest, "changes": changes}, output, root
    )


def replace_part(path: Path, name: str, transform) -> None:
    with zipfile.ZipFile(path) as archive:
        parts = {item.filename: archive.read(item) for item in archive.infolist()}
    parts[name] = transform(parts.get(name, b""))
    with zipfile.ZipFile(path, "w") as archive:
        for item, data in parts.items():
            archive.writestr(item, data)


def test_greek_create_inspect_edit_preserves_original_styles_tables_and_formulas(
    tmp_path, monkeypatch
):
    def no_network(*args, **kwargs):
        raise AssertionError("Spreadsheet adapter must remain local")

    monkeypatch.setattr(socket.socket, "connect", no_network)
    source, result = create(tmp_path)
    original = source.read_bytes()
    assert result["formula_count"] == 1
    assert result["needs_recalculation"] is True
    assert result["formula_results"] == "not_calculated"
    assert result["cached_values"] == "not_read_or_trusted"
    assert result["visual_review"] == "not_performed"
    assert result["editable"] is True
    output, changed = edit(
        tmp_path,
        source,
        result["source_sha256"],
        [
            {"sheet": "Έσοδα", "cell": "B2", "value": 195.5},
            {"sheet": "Έσοδα", "cell": "B4", "formula": "=SUM(B2:B3)", "format": {"italic": True}},
        ],
    )
    assert source.read_bytes() == original
    assert changed["changes_applied"] == 2
    assert changed["original_sha256"] == hashlib.sha256(original).hexdigest()
    with closing(openpyxl.load_workbook(output)) as book:
        ws = book["Έσοδα"]
        assert ws["A3"].value == "Θεσσαλονίκη"
        assert ws["B2"].value == 195.5
        assert ws["B4"].value == "=SUM(B2:B3)"
        assert ws["B4"].number_format == '#,##0.00 "€"'
        assert ws["B4"].font.bold and ws["B4"].font.italic
        assert ws.freeze_panes == "A2"
        assert ws.tables["Revenue"].ref == "A1:B3"
        assert book.calculation.fullCalcOnLoad and book.calculation.forceFullCalc
    json.dumps(changed, ensure_ascii=False, allow_nan=False)


def test_only_explicit_formula_objects_make_formulas_or_links(tmp_path):
    literals = [
        "=1+1",
        "+1+1",
        "-1+1",
        "@SUM(A1:A2)",
        "https://example.com",
        "mailto:a@example.com",
        "#N/A",
    ]
    source, result = create(
        tmp_path, {"sheets": [{"name": "Data", "rows": [literals + [{"formula": "=1+1"}]]}]}
    )
    with closing(openpyxl.load_workbook(source)) as book:
        for column, literal in enumerate(literals, 1):
            cell = book.active.cell(1, column)
            assert cell.value == literal
            assert cell.data_type == "s"
            assert cell.hyperlink is None
        assert book.active["H1"].data_type == "f"
    output, edited = edit(
        tmp_path,
        source,
        result["source_sha256"],
        [
            {"sheet": "Data", "cell": "H1", "value": '=HYPERLINK("https://example.com")'},
            {"sheet": "Data", "cell": "A2", "value": "https://example.com"},
        ],
    )
    with closing(openpyxl.load_workbook(output)) as book:
        assert book.active["H1"].data_type == "s"
        assert book.active["A2"].hyperlink is None
    assert edited["formula_count"] == 0 and not edited["needs_recalculation"]


def test_formulas_never_return_cached_results_as_calculated(tmp_path):
    source = tmp_path / "source.xlsx"
    with xlsxwriter.Workbook(source) as book:
        book.add_worksheet("Data").write_formula("A1", "=2+2", None, 999)
    result = sheets.inspect(source)
    assert result["sheets"][0]["cells"][0]["formula"] == "=2+2"
    assert "999" not in json.dumps(result)
    assert result["needs_recalculation"]
    output, _ = edit(
        tmp_path, source, result["source_sha256"], [{"sheet": "Data", "cell": "A2", "value": 3}]
    )
    with closing(openpyxl.load_workbook(output, data_only=True)) as book:
        assert book.active["A1"].value is None


@pytest.mark.parametrize(
    "formula",
    [
        '=HYPERLINK("https://example.com")',
        '=WEBSERVICE("https://example.com")',
        "=cmd|' /C calc'!A1",
        "='[other.xlsx]Sheet1'!A1",
        "=Revenue[Amount]",
        "=NamedRange",
        "=SUM(A1:A10001)",
        "={1,2}",
        "=#REF!",
        "=",
        "SUM(A1:A2)",
        "='Missing'!A1",
    ],
)
def test_unsafe_or_unsupported_explicit_formulas_fail_before_write(tmp_path, formula):
    with pytest.raises(ValueError):
        create(tmp_path, {"sheets": [{"name": "Data", "rows": [[{"formula": formula}]]}]})
    assert not (tmp_path / "source.xlsx").exists()


def test_local_cross_sheet_formula_works_with_greek_and_quoted_apostrophe(tmp_path):
    source, result = create(
        tmp_path,
        {
            "sheets": [
                {"name": "Πωλήσεις", "rows": [[4], [6]]},
                {"name": "O'Brien", "rows": [[9]]},
                {"name": "Σύνοψη", "rows": [[{"formula": "=SUM('Πωλήσεις'!A1:A2)+'O''Brien'!A1"}]]},
            ]
        },
    )
    assert result["formula_count"] == 1 and result["editable"]


def test_stale_hash_and_original_alias_are_rejected(tmp_path):
    source, result = create(tmp_path)
    original = source.read_bytes()
    change = [{"sheet": "Έσοδα", "cell": "B2", "value": 10}]
    with pytest.raises(ValueError, match="stale"):
        edit(tmp_path, source, "0" * 64, change)
    with pytest.raises(ValueError, match="expected_source_sha256"):
        sheets.edit(source, {"changes": change}, tmp_path / "other.xlsx", tmp_path)
    with pytest.raises(ValueError, match="separate output"):
        sheets.edit(
            source,
            {"expected_source_sha256": result["source_sha256"], "changes": change},
            source,
            tmp_path,
        )
    alias = tmp_path / "alias.xlsx"
    alias.hardlink_to(source)
    with pytest.raises(ValueError, match="separate output"):
        sheets.edit(
            source,
            {"expected_source_sha256": result["source_sha256"], "changes": change},
            alias,
            tmp_path,
        )
    assert original == source.read_bytes()
    assert not (tmp_path / "revised.xlsx").exists()


def test_source_changed_during_validation_fails_without_output(tmp_path, monkeypatch):
    source, result = create(tmp_path)
    load = sheets._load

    def change_after_read(data):
        source.write_bytes(source.read_bytes() + b"changed")
        return load(data)

    monkeypatch.setattr(sheets, "_load", change_after_read)
    with pytest.raises(ValueError, match="changed during edit"):
        edit(
            tmp_path,
            source,
            result["source_sha256"],
            [{"sheet": "Έσοδα", "cell": "B2", "value": 10}],
        )
    assert not (tmp_path / "revised.xlsx").exists()


@pytest.mark.parametrize(
    "part",
    [
        "xl/vbaProject.bin",
        "xl/drawings/drawing1.xml",
        "xl/pivotTables/pivotTable1.xml",
        "customXml/item1.xml",
    ],
)
def test_unsupported_parts_reject_edit_without_losing_source(tmp_path, part):
    source, _ = create(tmp_path)
    replace_part(source, part, lambda _: b"<unsupported />" if part.endswith(".xml") else b"macro")
    original = source.read_bytes()
    with pytest.raises(ValueError, match="Unsupported package part"):
        edit(
            tmp_path,
            source,
            hashlib.sha256(original).hexdigest(),
            [{"sheet": "Έσοδα", "cell": "B2", "value": 10}],
        )
    assert original == source.read_bytes()
    assert not (tmp_path / "revised.xlsx").exists()


@pytest.mark.parametrize(
    "feature",
    [
        "<extLst />",
        "<conditionalFormatting />",
        "<dataValidations />",
        "<sheetProtection />",
        "<ignoredErrors />",
        "<unknownNewFeature />",
    ],
)
def test_unsupported_sheet_features_reject_edit(tmp_path, feature):
    source, _ = create(tmp_path)
    replace_part(
        source,
        "xl/worksheets/sheet1.xml",
        lambda raw: raw.replace(b"</worksheet>", feature.encode() + b"</worksheet>"),
    )
    with pytest.raises(ValueError, match="safely edited"):
        edit(
            tmp_path,
            source,
            hashlib.sha256(source.read_bytes()).hexdigest(),
            [{"sheet": "Έσοδα", "cell": "B2", "value": 10}],
        )


def test_external_relationship_is_reported_and_rejected(tmp_path):
    source = tmp_path / "source.xlsx"
    with xlsxwriter.Workbook(source) as book:
        book.add_worksheet("Data").write_url("A1", "https://example.com")
    result = sheets.inspect(source)
    assert not result["editable"]
    assert "External relationships are unsupported" in result["edit_blockers"]
    with pytest.raises(ValueError, match="External relationships"):
        edit(
            tmp_path,
            source,
            result["source_sha256"],
            [{"sheet": "Data", "cell": "B2", "value": 10}],
        )


def test_macro_extension_rejected(tmp_path):
    source, _ = create(tmp_path)
    macro = source.with_suffix(".xlsm")
    source.rename(macro)
    with pytest.raises(ValueError, match="macro-enabled"):
        sheets.inspect(macro)


@pytest.mark.parametrize(
    "spec",
    [
        {"sheets": []},
        {"sheets": [{"name": "Data", "rows": [[1]]}] * 17},
        {"sheets": [{"name": "Data"}, {"name": "data"}]},
        {"sheets": [{"name": "bad/name"}]},
        {"sheets": [{"name": "Data", "cells": {"IW1": 1}}]},
        {"sheets": [{"name": "Data", "cells": {"A10001": 1}}]},
        {"sheets": [{"name": "Data", "cells": {"IV10000": 1}}]},
        {"sheets": [{"name": "Data", "rows": [[float("nan")]]}]},
        {"sheets": [{"name": "Data", "rows": [[10**16]]}]},
        {"sheets": [{"name": "Data", "rows": [["a" * 32768]]}]},
        {"sheets": [{"name": "Data", "rows": [["a\x00b"]]}]},
        {"sheets": [{"name": "Data", "rows": [[{"value": 1, "formula": "=1"}]]}]},
        {"sheets": [{"name": "Data", "rows": [[{"value": 1, "format": {"align": []}}]]}]},
        {
            "sheets": [
                {
                    "name": "Data",
                    "rows": [[{"value": 1, "format": {"hyperlink": "https://example.com"}}]],
                }
            ]
        },
        {"sheets": [{"name": "Data", "rows": [[1]], "cells": {"A1": 2}}]},
        {"sheets": [{"name": "Data", "rows": [[1]], "column_widths": [float("nan")]}]},
        {"sheets": [{"name": "Data", "rows": [[1]], "unsupported": True}]},
    ],
)
def test_invalid_create_specs_fail_explicitly_without_output(tmp_path, spec):
    with pytest.raises(ValueError):
        create(tmp_path, spec)
    assert not (tmp_path / "source.xlsx").exists()


def test_package_dimension_and_merge_bombs_rejected_before_loader(tmp_path, monkeypatch):
    source, _ = create(tmp_path)
    replace_part(
        source,
        "xl/worksheets/sheet1.xml",
        lambda raw: raw.replace(b'ref="A1:B4"', b'ref="A1:IV10000"'),
    )
    monkeypatch.setattr(sheets, "_load", lambda *_: pytest.fail("must not reach loader"))
    with pytest.raises(ValueError, match="limits"):
        sheets.inspect(source)


def test_package_entity_and_duplicate_parts_rejected(tmp_path):
    source, _ = create(tmp_path)
    replace_part(
        source,
        "xl/worksheets/sheet1.xml",
        lambda raw: b'<!DOCTYPE worksheet [<!ENTITY bomb "unsafe">]>' + raw,
    )
    with pytest.raises(ValueError, match="DTDs"):
        sheets.inspect(source)


def test_duplicate_zip_parts_rejected(tmp_path):
    source, _ = create(tmp_path)
    with zipfile.ZipFile(source, "a") as archive, pytest.warns(UserWarning):
        archive.writestr("xl/workbook.xml", b"<workbook />")
    with pytest.raises(ValueError, match="duplicate package"):
        sheets.inspect(source)


def test_repeated_merge_expansion_rejected_before_loader(tmp_path, monkeypatch):
    source, _ = create(tmp_path)
    merges = b"<mergeCells>" + b'<mergeCell ref="A1:E10000"/>' * 2 + b"</mergeCells>"
    replace_part(
        source,
        "xl/worksheets/sheet1.xml",
        lambda raw: raw.replace(b"</worksheet>", merges + b"</worksheet>"),
    )
    monkeypatch.setattr(sheets, "_load", lambda *_: pytest.fail("must not reach loader"))
    with pytest.raises(ValueError, match="merged-cell area"):
        sheets.inspect(source)


def test_empty_create_and_excel_future_function_roundtrip(tmp_path):
    source, result = create(
        tmp_path,
        {
            "sheets": [
                {"name": "Empty"},
                {"name": "Data", "rows": [[{"formula": '=CONCAT("a","b")'}]]},
            ]
        },
    )
    assert result["editable"] and result["sheet_count"] == 2
    output, result = edit(
        tmp_path,
        source,
        result["source_sha256"],
        [{"sheet": "Empty", "cell": "A1", "value": "filled"}],
    )
    assert result["editable"] and result["formula_count"] == 1


def test_inspect_non_finite_values_stays_json_safe_and_edit_rejects(tmp_path):
    source, _ = create(tmp_path, {"sheets": [{"name": "Data", "rows": [[1]]}]})
    replace_part(
        source, "xl/worksheets/sheet1.xml", lambda raw: raw.replace(b"<v>1</v>", b"<v>1e999</v>")
    )
    result = sheets.inspect(source)
    json.dumps(result, allow_nan=False)
    assert not result["editable"]
    with pytest.raises(ValueError, match="Non-finite"):
        edit(
            tmp_path, source, result["source_sha256"], [{"sheet": "Data", "cell": "A1", "value": 1}]
        )


def test_edit_table_header_and_non_anchor_merge_rejected(tmp_path):
    source, result = create(tmp_path)
    with pytest.raises(ValueError, match="table header"):
        edit(
            tmp_path,
            source,
            result["source_sha256"],
            [{"sheet": "Έσοδα", "cell": "A1", "value": "changed"}],
        )
    with closing(openpyxl.load_workbook(source)) as book:
        book.active.merge_cells("A6:B6")
        book.save(source)
    digest = sheets.inspect(source)["source_sha256"]
    with pytest.raises(ValueError, match="merged cell"):
        edit(tmp_path, source, digest, [{"sheet": "Έσοδα", "cell": "B6", "value": "changed"}])


def test_inspection_is_bounded_and_json_serializable_for_dates(tmp_path):
    source = tmp_path / "source.xlsx"
    book = openpyxl.Workbook()
    book.active["A1"] = datetime(2026, 10, 6, 14, 0)
    for row in range(2, 1200):
        book.active.cell(row, 1, "Στοιχεία")
    book.save(source)
    result = sheets.inspect(source)
    assert result["cell_count"] == 1199
    assert len(result["sheets"][0]["cells"]) == sheets.MAX_PREVIEW_CELLS
    assert result["sheets"][0]["cells_truncated"]
    assert result["sheets"][0]["cells"][0]["value"] == "2026-10-06T14:00:00"
    json.dumps(result, allow_nan=False)


def test_formatting_and_literal_table_headers(tmp_path):
    source, result = create(
        tmp_path,
        {
            "sheets": [
                {
                    "name": "Data",
                    "rows": [
                        ["=Literal", "https://example.com"],
                        [
                            {
                                "value": "00123",
                                "format": {
                                    "fill": "#ffeeaa",
                                    "font_color": "#123456",
                                    "align": "center",
                                    "wrap": True,
                                },
                            },
                            True,
                        ],
                    ],
                    "tables": [{"range": "A1:B2", "name": "SafeTable"}],
                }
            ]
        },
    )
    with closing(openpyxl.load_workbook(source)) as book:
        assert book.active["A1"].data_type == "s"
        assert book.active["B1"].hyperlink is None
        cell = book.active["A2"]
        assert cell.value == "00123"
        assert cell.fill.fgColor.rgb == "FFFFEEAA"
        assert cell.font.color.rgb == "FF123456"
        assert cell.alignment.horizontal == "center" and cell.alignment.wrap_text
        assert book.active["B2"].value is True
    assert result["editable"] and not result["needs_recalculation"]
