"""Independent content checks for skill evals; visual quality still needs review."""

import csv
import hashlib
import json
import re
import sys
from decimal import Decimal
from pathlib import Path


def unchanged():
    for name, expected in json.loads(Path("source-hashes.json").read_text()).items():
        assert hashlib.sha256(Path(name).read_bytes()).hexdigest() == expected, name


def check():
    unchanged()
    kind = sys.argv[1]
    if kind == "unchanged":
        assert not any(Path(".").glob("report.*"))
        assert not Path("comparison.md").exists()
        return
    if kind == "research":
        text = Path("comparison.md").read_text()
        assert "policy-a.md" in text and "policy-b.md" in text
        # Budget checks tolerate normal digit grouping. This checks source
        # coverage, not the quality of the model's recommendations.
        numbers = re.sub(r"[,\u00a0 ]", "", text)
        assert "25000" in numbers and "28000" in numbers
        return
    with Path("services.csv").open(newline="") as source:
        rows = list(csv.DictReader(source))
    totals = [sum(Decimal(row[column]) for row in rows) for column in ("orders", "revenue", "cost")]
    totals.append(totals[1] - totals[2])
    if kind == "xlsx":
        from openpyxl import load_workbook

        book = load_workbook("report.xlsx")
        cached = load_workbook("report.xlsx", data_only=True)
        try:
            data = list(book["Data"].values)
            assert list(data[0]) == list(rows[0])
            for actual, expected in zip(data[1:], rows, strict=True):
                assert actual[0] == expected["service"]
                assert [Decimal(str(value)) for value in actual[1:]] == [
                    Decimal(expected[key]) for key in ("orders", "revenue", "cost")
                ]
            for row, expected in enumerate(totals, 2):
                assert book["Summary"].cell(row, 2).data_type == "f"
                assert Decimal(str(cached["Summary"].cell(row, 2).value)) == expected
        finally:
            book.close()
            cached.close()
        return
    if kind == "docx":
        from docx import Document

        doc = Document("report.docx")
        assert doc.tables, "The requested table must remain editable"
        text = "\n".join(
            [paragraph.text for paragraph in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )
    elif kind == "pptx":
        from pptx import Presentation

        deck = Presentation("report.pptx")
        assert len(deck.slides) == 3
        text = "\n".join(
            shape.text
            if shape.has_text_frame
            else "\n".join(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.has_table
            else ""
            for slide in deck.slides
            for shape in slide.shapes
        )
    elif kind == "pdf":
        import fitz

        with fitz.open("report.pdf") as doc:
            assert len(doc) == 1
            text = "\n".join(page.get_text() for page in doc)
    else:
        raise ValueError(f"Unsupported check: {kind}")
    assert all(row["service"] in text for row in rows)
    assert all(str(total) in text for total in totals)
    assert "overhead" in text.casefold()


if __name__ == "__main__":
    check()
