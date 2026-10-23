"""Bounded DOCX authoring and selected text edits, without a whole-package resave.

The command layer owns path authorization, staging, and atomic publication. This
adapter writes only its supplied, new output file after validating all edits.
Narrow package-preserving editing is adapted from the app's work-documents helper.
"""

from __future__ import annotations

import hashlib
import re
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
DOC_PART = "word/document.xml"
MAIN_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
MAX_SOURCE_BYTES = 50 * 1024 * 1024
MAX_UNPACKED = 100 * 1024 * 1024
MAX_PARAGRAPHS = 10_000
MAX_CHARACTERS = 1_000_000

GUIDE = {
    "format": "docx",
    "create": {
        "title": "Weekly brief",
        "blocks": [
            {"kind": "heading", "text": "Overview", "level": 1},
            {"kind": "paragraph", "text": "Draft text. Ελληνικό κείμενο."},
            {"kind": "table", "rows": [["Item", "Value"], ["Total", "10"]]},
        ],
    },
    "edit": {
        "expected_source_sha256": "Copy source_sha256 from inspect",
        "edits": [
            {
                "locator": "paragraph:2",
                "expected_text_sha256": "Copy this paragraph's text_sha256 from inspect",
                "find": "Overview",
                "replace": "Updated overview",
            }
        ],
    },
    "limits": [
        "Create: 1-200 blocks; heading levels 1-9; tables 1-200 rows and 1-8 columns.",
        "The optional title is inserted as the first paragraph.",
        "Inspect: main-body/table paragraphs only; headers, footers, notes and images are not extracted.",
        "Edit: 1-100 inline replacements, one per editable paragraph, each find occurs exactly once.",
        "Edits require both source and paragraph SHA-256; unsupported complex selections fail.",
        "Replacement inherits formatting at the match start; pagination can change.",
        "DOCX only; no macros, embedded files, active controls, signatures or external links other than hyperlinks.",
        "No docxtpl execution, renderer or visual approval in this adapter.",
    ],
}


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _text(value: object, name: str, *, limit: int = 100_000, empty: bool = False) -> str:
    if not isinstance(value, str) or len(value) > limit or (not empty and not value.strip()):
        raise ValueError(f"{name} must be text with at most {limit} characters")
    if any(
        (ord(char) < 32 and char not in "\r\n\t")
        or 0xD800 <= ord(char) <= 0xDFFF
        or ord(char) in (0xFFFE, 0xFFFF)
        for char in value
    ):
        raise ValueError(f"{name} contains characters unsupported by document XML")
    return value


def _object(value: object, allowed: set[str], name: str) -> dict:
    if not isinstance(value, dict):
        raise ValueError(f"{name} must be an object")
    if value.keys() - allowed:
        raise ValueError(f"Unsupported {name} fields: {', '.join(sorted(value.keys() - allowed))}")
    return value


def _spec(spec: dict, allowed: set[str]) -> None:
    _object(spec, allowed | {"schema_version"}, "specification")
    if "schema_version" in spec and (
        type(spec["schema_version"]) is not int or spec["schema_version"] != 1
    ):
        raise ValueError("Only schema_version 1 is supported")


def _xml(data: bytes):
    try:
        root = etree.fromstring(
            data,
            etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False),
        )
    except etree.XMLSyntaxError as error:
        raise ValueError("Invalid DOCX XML") from error
    # Checking the parsed declaration also covers UTF-16 XML encodings.
    if root.getroottree().docinfo.doctype:
        raise ValueError("DOCX XML document type declarations are unsupported")
    return root


def _read(source: Path) -> bytes:
    with source.open("rb") as stream:
        data = stream.read(MAX_SOURCE_BYTES + 1)
    if len(data) > MAX_SOURCE_BYTES:
        raise ValueError("DOCX source exceeds 50 MiB")
    return data


def _parts(data: bytes) -> dict[str, bytes]:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_UNPACKED:
                raise ValueError("DOCX package exceeds expanded size limits")
            names = [item.filename for item in entries]
            if len(set(names)) != len(names):
                raise ValueError("Duplicate DOCX package entries")
            if any(
                "\\" in name or name.startswith("/") or ".." in name.split("/") for name in names
            ):
                raise ValueError("Invalid DOCX package entry path")
            if any(item.flag_bits & 1 for item in entries):
                raise ValueError("Encrypted DOCX packages are unsupported")
            parts = {item.filename: archive.read(item) for item in entries}
    except (BadZipFile, RuntimeError, NotImplementedError) as error:
        raise ValueError("File is not a readable DOCX package") from error
    if not {DOC_PART, "[Content_Types].xml", "_rels/.rels"}.issubset(parts):
        raise ValueError("Required DOCX package parts are missing")
    for name, value in parts.items():
        if any(
            marker in name.lower()
            for marker in ("vbaproject", "embeddings/", "activex/", "_xmlsignatures/")
        ):
            raise ValueError(
                "Macros, active controls, embedded files and signed DOCX packages are unsupported"
            )
        if name.lower().endswith((".xml", ".rels")):
            root = _xml(value)
            if name.lower().endswith(".rels"):
                for relation in root:
                    if relation.get("TargetMode", "").lower() == "external" and not relation.get(
                        "Type", ""
                    ).endswith("/hyperlink"):
                        raise ValueError("Linked external content is unsupported")
    content_types = _xml(parts["[Content_Types].xml"])
    if not any(
        item.get("PartName") == "/" + DOC_PART and item.get("ContentType") == MAIN_TYPE
        for item in content_types
    ):
        raise ValueError("Only standard, non-macro DOCX documents are supported")
    relations = _xml(parts["_rels/.rels"])
    main = [item for item in relations if item.get("Type", "").endswith("/officeDocument")]
    if len(main) != 1 or main[0].get("Target", "").lstrip("/") != DOC_PART:
        raise ValueError("Unsupported DOCX main document relationship")
    root = _xml(parts[DOC_PART])
    if root.tag != qn("w:document") or root.find(qn("w:body")) is None:
        raise ValueError("DOCX main part must contain a Word document body")
    return parts


def _paragraphs(root) -> list:
    paragraphs = root.xpath("/w:document/w:body//w:p", namespaces=NS)
    if len(paragraphs) > MAX_PARAGRAPHS:
        raise ValueError("DOCX exceeds 10000 main-body paragraphs")
    return paragraphs


def _paragraph_text(paragraph) -> str:
    values = []
    for node in paragraph.iter():
        # Do not duplicate text from a nested text-box paragraph in its parent.
        parent = next((p for p in node.iterancestors() if p.tag == qn("w:p")), None)
        if parent is not paragraph:
            continue
        if node.tag == qn("w:t"):
            values.append(node.text or "")
        elif node.tag == qn("w:tab"):
            values.append("\t")
        elif node.tag in (qn("w:br"), qn("w:cr")):
            values.append("\n")
    return "".join(values)


def _document_edit_restriction(parts: dict, root) -> str | None:
    if root.xpath(
        "//w:ins | //w:del | //w:moveFrom | //w:moveTo | //w:pPrChange | //w:rPrChange | //w:tblPrChange | //w:trPrChange | //w:tcPrChange | //w:sectPrChange | //w:altChunk",
        namespaces=NS,
    ):
        return "Tracked changes and imported content require a separate editing workflow"
    if "word/settings.xml" in parts:
        settings = _xml(parts["word/settings.xml"])
        for protection in settings.iter(qn("w:documentProtection")):
            if protection.get(qn("w:enforcement"), "0").lower() in ("1", "true", "on"):
                return "Document protection is enabled"
    return None


def _plain_nodes(paragraph) -> list:
    for ancestor in paragraph.iterancestors():
        if ancestor.tag not in {
            qn("w:document"),
            qn("w:body"),
            qn("w:tbl"),
            qn("w:tr"),
            qn("w:tc"),
        }:
            raise ValueError("Selected paragraph is inside unsupported complex content")
    nodes = []
    for child in paragraph:
        if child.tag == qn("w:pPr"):
            continue
        if child.tag != qn("w:r"):
            raise ValueError("Selected paragraph has complex content")
        for node in child:
            if node.tag == qn("w:rPr"):
                continue
            if node.tag != qn("w:t"):
                raise ValueError("Selected paragraph has non-text run content")
            nodes.append(node)
    return nodes


def _metadata(data: bytes, parts: dict | None = None) -> dict:
    parts = _parts(data) if parts is None else parts
    root = _xml(parts[DOC_PART])
    restriction = _document_edit_restriction(parts, root)
    records = []
    characters = 0
    for index, paragraph in enumerate(_paragraphs(root), 1):
        text = _paragraph_text(paragraph)
        characters += len(text)
        if characters > MAX_CHARACTERS:
            raise ValueError("DOCX main-body text exceeds 1000000 characters")
        reason = restriction
        if reason is None:
            try:
                _plain_nodes(paragraph)
            except ValueError as error:
                reason = str(error)
        records.append(
            {
                "locator": f"paragraph:{index}",
                "text": text,
                "text_sha256": _digest(text.encode("utf-8")),
                "editable": reason is None,
                "edit_restriction": reason,
            }
        )
    warnings = [
        "Extraction covers main-body and table text only; headers, footers, notes, images and layout are not verified.",
        "Page count and visual review require a separate renderer.",
    ]
    complex_content = root.xpath(
        "//w:drawing | //w:pict | //w:object | //w:altChunk | //w:del | //w:instrText | //w:fldSimple | //w:sym",
        namespaces=NS,
    )
    restricted_paragraphs = sum(not record["editable"] for record in records)
    if complex_content or restricted_paragraphs:
        warnings.append(
            "Complex main-body content is only partially extracted; field values are not recalculated."
        )
    if restriction:
        warnings.append(restriction)
    return {
        "format": "docx",
        "source_sha256": _digest(data),
        "paragraphs": records,
        "paragraph_count": len(records),
        "table_count": len(root.xpath("/w:document/w:body//w:tbl", namespaces=NS)),
        "package_part_count": len(parts),
        "coverage": {
            "status": "partial"
            if complex_content or restriction or restricted_paragraphs
            else "complete_for_supported_content",
            "scope": "main-body/table paragraph text; excludes headers, footers, notes, images and layout",
            "characters": characters,
            "editable_paragraphs": len(records) - restricted_paragraphs,
        },
        "warnings": warnings,
        "visual_review": "not_performed",
    }


def inspect(source: Path) -> dict:
    """Inspect a bounded DOCX and return one-based, source-bound paragraph choices."""
    return _metadata(_read(source))


def create(spec: dict, output: Path, workspace: Path) -> dict:
    """Create an editable, text-led DOCX at a new staged output path."""
    _spec(spec, {"title", "blocks"})
    title = _text(spec["title"], "title", limit=200) if "title" in spec else None
    blocks = spec.get("blocks")
    if not isinstance(blocks, list) or not 1 <= len(blocks) <= 200:
        raise ValueError("Provide between 1 and 200 content blocks")
    # Validate the entire specification before constructing potentially large tables.
    characters = len(title or "")
    paragraph_count = int(title is not None)
    for block in blocks:
        if not isinstance(block, dict):
            raise ValueError("Each content block must be an object")
        kind = block.get("kind")
        if kind in ("paragraph", "heading"):
            paragraph_count += 1
            _object(
                block, {"kind", "text", "level"} if kind == "heading" else {"kind", "text"}, "block"
            )
            characters += len(_text(block.get("text"), "block text", empty=kind == "paragraph"))
            if kind == "heading" and (
                type(block.get("level", 1)) is not int or not 1 <= block.get("level", 1) <= 9
            ):
                raise ValueError("Heading level must be an integer from 1 through 9")
        elif kind == "table":
            _object(block, {"kind", "rows"}, "table block")
            rows = block.get("rows")
            if (
                not isinstance(rows, list)
                or not 1 <= len(rows) <= 200
                or not isinstance(rows[0], list)
                or not 1 <= len(rows[0]) <= 8
            ):
                raise ValueError("A table needs 1-200 rows and 1-8 columns")
            width = len(rows[0])
            paragraph_count += len(rows) * width
            for row in rows:
                if not isinstance(row, list) or len(row) != width:
                    raise ValueError("All table rows must have the same column count")
                characters += sum(len(_text(cell, "cell", limit=5000, empty=True)) for cell in row)
        else:
            raise ValueError(f"Unsupported block kind: {kind}")
        if characters > MAX_CHARACTERS:
            raise ValueError("DOCX creation text exceeds 1000000 characters")
        if paragraph_count > MAX_PARAGRAPHS:
            raise ValueError("DOCX creation exceeds 10000 main-body paragraphs")
    document = Document()
    document.core_properties.title = title or ""
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ("Normal", "Title", *(f"Heading {i}" for i in range(1, 10))):
        document.styles[name].font.name = "Arial"
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(8)
    if title:
        document.add_paragraph(title, "Title")
    for block in blocks:
        if block["kind"] == "heading":
            document.add_heading(block["text"], level=block.get("level", 1))
        elif block["kind"] == "paragraph":
            document.add_paragraph(block["text"])
        else:
            rows = block["rows"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for row_index, values in enumerate(rows):
                for cell, value in zip(table.rows[row_index].cells, values, strict=True):
                    cell.text = value
                    if row_index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    stream = BytesIO()
    document.save(stream)
    data = stream.getvalue()
    result = _metadata(data)
    with output.open("xb") as staged:
        staged.write(data)
    return {**result, "operation": "create", "output_sha256": result["source_sha256"]}


def edit(source: Path, spec: dict, output: Path, workspace: Path) -> dict:
    """Apply guarded inline edits while retaining all other package-part bytes."""
    _spec(spec, {"expected_source_sha256", "edits"})
    data = _read(source)
    if spec.get("expected_source_sha256") != _digest(data):
        raise ValueError("Stale source: expected_source_sha256 must match the inspected DOCX")
    before = _parts(data)
    root = _xml(before[DOC_PART])
    restriction = _document_edit_restriction(before, root)
    if restriction:
        raise ValueError(restriction)
    paragraphs = _paragraphs(root)
    edits = spec.get("edits")
    if not isinstance(edits, list) or not 1 <= len(edits) <= 100:
        raise ValueError("Provide between 1 and 100 text edits")
    seen = set()
    changed = []
    expected_texts = {}
    for selection in edits:
        _object(selection, {"locator", "expected_text_sha256", "find", "replace"}, "edit")
        locator = _text(selection.get("locator"), "locator", limit=80)
        if not re.fullmatch(r"paragraph:[1-9][0-9]*", locator):
            raise ValueError("An edit must identify a one-based paragraph:N locator")
        if locator in seen:
            raise ValueError("Use one edit per paragraph in a revision")
        seen.add(locator)
        index = int(locator[10:]) - 1
        if index >= len(paragraphs):
            raise ValueError("Selected paragraph is unavailable")
        nodes = _plain_nodes(paragraphs[index])
        original = "".join(node.text or "" for node in nodes)
        if selection.get("expected_text_sha256") != _digest(original.encode("utf-8")):
            raise ValueError("Stale selection: paragraph differs from its expected_text_sha256")
        find = _text(selection.get("find"), "find")
        replacement = _text(selection.get("replace"), "replace", empty=True)
        if any(char in replacement for char in "\r\n\t"):
            raise ValueError("Edits support inline text replacements only")
        if original.count(find) != 1:
            raise ValueError("Find text must match exactly once in the selected paragraph")
        start = original.index(find)
        end = start + len(find)
        position = 0
        for node in nodes:
            value = node.text or ""
            next_position = position + len(value)
            if position < end and next_position > start:
                prefix = value[: max(0, start - position)]
                suffix = value[max(0, end - position) :]
                insert = replacement if position <= start < next_position else ""
                node.text = prefix + insert + suffix
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            position = next_position
        changed.append(locator)
        expected_texts[locator] = original[:start] + replacement + original[end:]
    updated_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    stream = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(stream, "w") as revised:
        revised.comment = original.comment
        for item in original.infolist():
            revised.writestr(
                item, updated_xml if item.filename == DOC_PART else before[item.filename]
            )
    result_data = stream.getvalue()
    after = _parts(result_data)
    if before.keys() != after.keys() or any(
        after[name] != value for name, value in before.items() if name != DOC_PART
    ):
        raise ValueError("Preservation failed: an unrelated package part changed")
    old_paragraphs = _paragraphs(_xml(before[DOC_PART]))
    new_paragraphs = _paragraphs(_xml(after[DOC_PART]))
    if len(old_paragraphs) != len(new_paragraphs):
        raise ValueError("Preservation failed: paragraph count changed")
    for index, (old, new) in enumerate(zip(old_paragraphs, new_paragraphs, strict=True), 1):
        locator = f"paragraph:{index}"
        if locator not in seen and etree.tostring(old) != etree.tostring(new):
            raise ValueError("Preservation failed: an unrelated paragraph changed")
        if locator in seen and _paragraph_text(new) != expected_texts[locator]:
            raise ValueError(
                "Edit verification failed: resulting text differs from the requested edit"
            )
    result = _metadata(result_data, after)
    with output.open("xb") as staged:
        staged.write(result_data)
    return {
        **result,
        "operation": "edit",
        "input_sha256": _digest(data),
        "output_sha256": result["source_sha256"],
        "preservation": {
            "changed_paragraphs": changed,
            "unchanged_package_parts": len(before) - 1,
            "package_parts_preserved": True,
            "unrelated_paragraphs_preserved": True,
            "replacement_formatting": "inherits the formatting at the match start",
            "pagination": "not_verified",
        },
    }
