"""Bounded PPTX authoring and selective text edits; this module does not render slides.

The common office command owns package/path validation and atomic publication. Output
paths here must be disposable staging paths. Edits copy the original package, replacing
only touched slide XML, so unrelated package members are not reserialized by python-pptx.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
import unicodedata
import warnings
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

MAX_SLIDES = 50
MAX_SHAPES = 100
MAX_TEXT = 20_000
MAX_TOTAL_TEXT = 200_000
MAX_IMAGE_BYTES = 20 * 1024 * 1024
MAX_IMAGE_PIXELS = 25_000_000
EMU_PER_INCH = 914400

GUIDE = {
    "format": "pptx",
    "create": {
        "title": "Quarterly review",
        "slides": [{"title": "Επισκόπηση", "body": ["First point", "Second point"]}],
    },
    "elements": [
        {
            "type": "text",
            "text": "Text",
            "x": 0.6,
            "y": 1.5,
            "width": 5,
            "height": 2,
            "font_size": 24,
        },
        {
            "type": "table",
            "rows": [["Region", "Sales"], ["Athens", "42"]],
            "x": 0.6,
            "y": 4,
            "width": 6,
            "height": 2,
        },
        {"type": "image", "path": "assets/photo.png", "x": 7, "y": 1.5, "width": 5, "height": 4},
    ],
    "edit": {
        "expected_source_sha256": "Use source_sha256 from inspect",
        "edits": [
            {
                "slide": 1,
                "shape_id": 2,
                "expected_text_sha256": "Use text_sha256 from inspect",
                "text": "Revised title",
            }
        ],
    },
    "locators": "slide, paragraph, run, row and column are 1-based; shape_id is the PPTX shape ID. Include both paragraph and run for precise run edits. Include both row and column for a table cell.",
    "limits": {
        "slides": MAX_SLIDES,
        "shapes_per_slide": MAX_SHAPES,
        "characters_per_text": MAX_TEXT,
        "characters_total": MAX_TOTAL_TEXT,
        "table_rows": 20,
        "table_columns": 10,
        "inspection": "At most 5000 table cells, 20000 paragraph/run nodes, 8 group levels; imported tables at most 100 rows and 100 columns.",
        "images": "Local workspace PNG/JPEG only, <=20 MiB, <=25 million pixels; contain-fit preserves aspect ratio.",
        "editing": "Whole-target replacement requires one paragraph and at most one plain run. Rich text requires a paragraph/run locator. Grouped shapes, fields, merged cells and non-text shapes cannot be edited.",
        "layout": "16:9 slides; inch geometry. Shape bounds are checked. Text-fit checks are conservative estimates, not rendered overflow verification. No rendering or full-fidelity guarantee.",
    },
}


def _hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _keys(value: object, allowed: set[str], label: str) -> dict:
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be an object")
    unknown = value.keys() - allowed
    if unknown:
        raise ValueError(f"Unsupported {label} fields: {', '.join(sorted(map(str, unknown)))}")
    return value


def _text(value: object, label: str) -> str:
    if not isinstance(value, str) or len(value) > MAX_TEXT:
        raise ValueError(f"{label} must be text with at most {MAX_TEXT} characters")
    if any(unicodedata.category(c) == "Cs" or (ord(c) < 32 and c not in "\n\t") for c in value):
        raise ValueError(f"{label} contains unsupported control characters")
    return value


def _number(value: object, label: str, low: float, high: float) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{label} must be a finite number")
    if not low <= value <= high:
        raise ValueError(f"{label} must be between {low} and {high}")
    return float(value)


def _integer(value: object, label: str, low: int, high: int) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or not low <= value <= high:
        raise ValueError(f"{label} must be an integer between {low} and {high}")
    return value


def _bool(value: object, label: str) -> bool:
    if not isinstance(value, bool):
        raise ValueError(f"{label} must be boolean")
    return value


def _bounds(element: dict, slide_width: int, slide_height: int) -> tuple[int, ...]:
    values = tuple(
        Inches(_number(element.get(k), k, 0 if k in ("x", "y") else 0.1, 56))
        for k in ("x", "y", "width", "height")
    )
    x, y, width, height = values
    if x + width > slide_width or y + height > slide_height:
        raise ValueError("Shape bounds extend outside the slide")
    return values


def _fit(text: str, width: int, height: int, font_size: float) -> None:
    """Reject likely overflow; this is deliberately not a font renderer."""
    width_pt = width / EMU_PER_INCH * 72 - 14.4
    height_pt = height / EMU_PER_INCH * 72 - 7.2
    if width_pt <= 0 or height_pt <= 0:
        raise ValueError("Text frame is too small")
    # A conservative width allowance includes Greek and wide Unicode characters.
    lines = 0
    for line in text.split("\n"):
        units = sum(
            1.0 if unicodedata.east_asian_width(c) in "WF" else 0.62 for c in line.expandtabs(4)
        )
        lines += max(1, math.ceil(units * font_size / width_pt))
    if lines * font_size * 1.2 > height_pt + 0.5:
        raise ValueError(
            "Text is likely to overflow its frame (estimated, not rendered); shorten it or enlarge the frame"
        )


def _write_text(
    frame, text: str, font_size: float, *, bold: bool = False, color: str = "243247"
) -> None:
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Inches(0.1)
    frame.margin_top = frame.margin_bottom = Inches(0.05)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = 1.2
        paragraph.space_before = paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def _image_bytes(value: object, workspace: Path) -> tuple[bytes, int, int]:
    if not isinstance(value, str) or not value or ":" in value or "\\" in value:
        raise ValueError("Image path must be a relative workspace path")
    relative = Path(value)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("Image path must be inside the workspace without traversal")
    root = workspace.resolve()
    candidate = root
    for part in relative.parts:
        candidate /= part
        if candidate.is_symlink():
            raise ValueError("Image path must not contain symbolic links")
    if not candidate.resolve().is_relative_to(root) or not candidate.is_file():
        raise ValueError("Image must be a regular file inside the workspace")
    info = candidate.stat()
    if info.st_nlink != 1 or info.st_size > MAX_IMAGE_BYTES:
        raise ValueError("Image must not be hard linked and must be at most 20 MiB")
    data = candidate.read_bytes()
    if len(data) > MAX_IMAGE_BYTES:
        raise ValueError("Image must be at most 20 MiB")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(data)) as image:
                width, height = image.size
                if image.format not in {"PNG", "JPEG"} or width * height > MAX_IMAGE_PIXELS:
                    raise ValueError("Image must be PNG/JPEG with at most 25 million pixels")
                image.verify()
            with Image.open(io.BytesIO(data)) as image:
                image.load()
    except (
        OSError,
        SyntaxError,
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
    ) as exc:
        raise ValueError("Image could not be safely decoded") from exc
    return data, width, height


def create(spec: dict, output: Path, workspace: Path) -> dict:
    """Create a bounded, editable deck at a caller-owned staging path."""
    _keys(spec, {"title", "slides"}, "presentation")
    slides = spec.get("slides")
    if not isinstance(slides, list) or not 1 <= len(slides) <= MAX_SLIDES:
        raise ValueError(f"slides must contain 1 to {MAX_SLIDES} slide objects")
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333333), Inches(7.5)
    title = _text(spec.get("title", ""), "title")
    if len(title) > 255:
        raise ValueError("Presentation metadata title must be at most 255 characters")
    presentation.core_properties.title = title
    total_text = 0
    for slide_spec in slides:
        _keys(slide_spec, {"title", "body", "elements"}, "slide")
        elements = slide_spec.get("elements", [])
        if not isinstance(elements, list):
            raise ValueError("elements must be a list")
        elements = list(elements)
        if "title" in slide_spec:
            elements.insert(
                0,
                {
                    "type": "text",
                    "text": slide_spec["title"],
                    "x": 0.5,
                    "y": 0.3,
                    "width": 12.3,
                    "height": 1.0,
                    "font_size": 32,
                    "bold": True,
                },
            )
        if "body" in slide_spec:
            body = slide_spec["body"]
            if isinstance(body, list):
                if len(body) > 100:
                    raise ValueError("body must have at most 100 paragraphs")
                body = "\n".join(_text(item, "body paragraph") for item in body)
            elements.insert(
                1 if "title" in slide_spec else 0,
                {
                    "type": "text",
                    "text": body,
                    "x": 0.6,
                    "y": 1.5,
                    "width": 12.1,
                    "height": 5.4,
                    "font_size": 24,
                },
            )
        if not 1 <= len(elements) <= MAX_SHAPES:
            raise ValueError(f"Each slide must contain 1 to {MAX_SHAPES} shapes")
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for element in elements:
            if not isinstance(element, dict):
                raise ValueError("Each element must be an object")
            kind = element.get("type")
            common = {"type", "x", "y", "width", "height"}
            allowed = {
                "text": {"text", "font_size", "bold", "color"},
                "table": {"rows", "font_size", "header"},
                "image": {"path"},
            }
            if not isinstance(kind, str) or kind not in allowed:
                raise ValueError("Supported element types are text, table and image")
            _keys(element, common | allowed[kind], "element")
            x, y, width, height = _bounds(
                element, presentation.slide_width, presentation.slide_height
            )
            if kind == "text":
                text = _text(element.get("text"), "text")
                size = _number(element.get("font_size", 24), "font_size", 10, 72)
                bold = _bool(element.get("bold", False), "bold")
                color = element.get("color", "243247")
                if not isinstance(color, str) or not re.fullmatch(r"[0-9A-Fa-f]{6}", color):
                    raise ValueError("color must be six hexadecimal RGB digits")
                _fit(text, width, height, size)
                shape = slide.shapes.add_textbox(x, y, width, height)
                _write_text(shape.text_frame, text, size, bold=bold, color=color)
                total_text += len(text)
            elif kind == "table":
                rows = element.get("rows")
                if (
                    not isinstance(rows, list)
                    or not 1 <= len(rows) <= 20
                    or not isinstance(rows[0], list)
                    or not 1 <= len(rows[0]) <= 10
                ):
                    raise ValueError("Table must have 1 to 20 rows and 1 to 10 columns")
                columns = len(rows[0])
                if any(not isinstance(row, list) or len(row) != columns for row in rows):
                    raise ValueError("Table rows must have equal column counts")
                size = _number(element.get("font_size", 16), "font_size", 10, 36)
                header = _bool(element.get("header", True), "header")
                table = slide.shapes.add_table(len(rows), columns, x, y, width, height).table
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        text = _text(value, "table cell")
                        _fit(text, width // columns, height // len(rows), size)
                        cell = table.cell(r, c)
                        _write_text(
                            cell.text_frame,
                            text,
                            size,
                            bold=header and r == 0,
                            color="FFFFFF" if header and r == 0 else "243247",
                        )
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(
                            "243247" if header and r == 0 else ("EEF2F6" if r % 2 else "FFFFFF")
                        )
                        total_text += len(text)
                table.first_row = header
            else:
                data, pixel_width, pixel_height = _image_bytes(element.get("path"), workspace)
                scale = min(width / pixel_width, height / pixel_height)
                image_width, image_height = round(pixel_width * scale), round(pixel_height * scale)
                slide.shapes.add_picture(
                    io.BytesIO(data),
                    x + (width - image_width) // 2,
                    y + (height - image_height) // 2,
                    image_width,
                    image_height,
                )
            if total_text > MAX_TOTAL_TEXT:
                raise ValueError(f"Presentation exceeds {MAX_TOTAL_TEXT} text characters")
    presentation.save(str(output))
    result = inspect(output)
    result["status"] = "created"
    return result


def _plain_frame(frame) -> bool:
    if len(frame.paragraphs) != 1:
        return False
    paragraph = frame.paragraphs[0]
    return len(paragraph.runs) <= 1 and all(
        child.tag.rsplit("}", 1)[-1] in {"pPr", "r", "endParaRPr"} for child in paragraph._p
    )


def _frame_info(frame) -> dict:
    return {
        "text": frame.text,
        "text_sha256": _hash(frame.text),
        "whole_text_editable": _plain_frame(frame),
        "paragraphs": [
            {
                "paragraph": p_index,
                "text": paragraph.text,
                "text_sha256": _hash(paragraph.text),
                "runs": [
                    {"run": r_index, "text": run.text, "text_sha256": _hash(run.text)}
                    for r_index, run in enumerate(paragraph.runs, 1)
                ],
            }
            for p_index, paragraph in enumerate(frame.paragraphs, 1)
        ],
    }


def _load(source: Path):
    try:
        presentation = Presentation(str(source))
        if not 1 <= len(presentation.slides) <= MAX_SLIDES:
            raise ValueError(f"Presentation must contain 1 to {MAX_SLIDES} slides")
        characters, text_nodes, cells = 0, 0, 0
        for slide in presentation.slides:
            queue = [(shape, 0) for shape in slide.shapes]
            count = 0
            while queue:
                shape, depth = queue.pop()
                count += 1
                if count > MAX_SHAPES or depth > 8:
                    raise ValueError("Presentation exceeds shape or group-depth limits")
                frames = [shape.text_frame] if shape.has_text_frame else []
                if shape.has_table:
                    table = shape.table
                    if len(table.rows) > 100 or len(table.columns) > 100:
                        raise ValueError("Presentation table exceeds inspection limits")
                    cells += len(table.rows) * len(table.columns)
                    if cells > 5000:
                        raise ValueError("Presentation exceeds 5000 table cells")
                    frames.extend(cell.text_frame for row in table.rows for cell in row.cells)
                for frame in frames:
                    characters += len(frame.text)
                    text_nodes += sum(1 + len(p.runs) for p in frame.paragraphs)
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    queue.extend((child, depth + 1) for child in shape.shapes)
                if characters > MAX_TOTAL_TEXT or text_nodes > 20_000:
                    raise ValueError("Presentation exceeds text inspection limits")
        return presentation
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Cannot read a supported PPTX presentation") from exc


def inspect(source: Path) -> dict:
    """Return package text and stable locators; no visual-verification claim."""
    presentation = _load(source)
    slides = []
    for index, slide in enumerate(presentation.slides, 1):
        shapes = []
        for shape in slide.shapes:
            item = {
                "shape_id": shape.shape_id,
                "name": shape.name,
                "type": str(shape.shape_type),
                "bounds_inches": {
                    "x": shape.left / EMU_PER_INCH,
                    "y": shape.top / EMU_PER_INCH,
                    "width": shape.width / EMU_PER_INCH,
                    "height": shape.height / EMU_PER_INCH,
                },
                "outside_slide": shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > presentation.slide_width
                or shape.top + shape.height > presentation.slide_height,
                "text_editable": bool(shape.has_text_frame),
            }
            if shape.has_text_frame:
                item.update(_frame_info(shape.text_frame))
            elif shape.has_table:
                item["cells"] = [
                    {
                        "row": r,
                        "column": c,
                        "merged": cell.is_merge_origin or cell.is_spanned,
                        **_frame_info(cell.text_frame),
                    }
                    for r, row in enumerate(shape.table.rows, 1)
                    for c, cell in enumerate(row.cells, 1)
                ]
            elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                item["unsupported_edit_reason"] = (
                    "Grouped shapes are preserved; text editing inside groups is unsupported"
                )
            else:
                item["unsupported_edit_reason"] = "This shape is not a supported text target"
            shapes.append(item)
        slides.append({"slide": index, "slide_id": slide.slide_id, "shapes": shapes})
    return {
        "format": "pptx",
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "slide_count": len(slides),
        "slides": slides,
        "visual_review": "not_performed",
        "text_overflow": "not_verified",
        "geometry_check": "unrotated_shape_bounds_only",
        "warnings": [
            "PPTX text/layout inspection is not rendered verification; use a renderer and visual review."
        ],
    }


def edit(source: Path, spec: dict, output: Path, workspace: Path) -> dict:
    """Apply hash-guarded text replacements without saving unrelated package parts."""
    del workspace  # No asset reads or external operations during edits.
    if source.resolve() == output.resolve() or (output.exists() and source.samefile(output)):
        raise ValueError("Editing requires a new output file; the original must be preserved")
    _keys(spec, {"expected_source_sha256", "edits"}, "edit specification")
    original = source.read_bytes()
    if spec.get("expected_source_sha256") != hashlib.sha256(original).hexdigest():
        raise ValueError("Stale source: expected_source_sha256 does not match")
    edits = spec.get("edits")
    if not isinstance(edits, list) or not 1 <= len(edits) <= 200:
        raise ValueError("edits must contain 1 to 200 replacements")
    presentation = _load(source)
    changed_parts = {}
    targets = set()
    for change in edits:
        _keys(
            change,
            {
                "slide",
                "shape_id",
                "row",
                "column",
                "paragraph",
                "run",
                "expected_text_sha256",
                "text",
            },
            "text edit",
        )
        slide_index = _integer(change.get("slide"), "slide", 1, len(presentation.slides))
        slide = presentation.slides[slide_index - 1]
        shape_id = _integer(change.get("shape_id"), "shape_id", 1, 2**32 - 1)
        shape = next((item for item in slide.shapes if item.shape_id == shape_id), None)
        if shape is None:
            raise ValueError("Unknown shape_id on the selected slide")
        text = _text(change.get("text"), "replacement text")
        if ("row" in change) != ("column" in change) or ("paragraph" in change) != (
            "run" in change
        ):
            raise ValueError("row/column and paragraph/run locators must be supplied in pairs")
        width, height = shape.width, shape.height
        if "row" in change:
            if not shape.has_table:
                raise ValueError("row/column requires a table")
            table = shape.table
            row = _integer(change["row"], "row", 1, len(table.rows))
            column = _integer(change["column"], "column", 1, len(table.columns))
            cell = table.cell(row - 1, column - 1)
            if cell.is_merge_origin or cell.is_spanned:
                raise ValueError("Merged table cells cannot be edited")
            frame = cell.text_frame
            width, height = table.columns[column - 1].width, table.rows[row - 1].height
        elif shape.has_text_frame:
            frame = shape.text_frame
        else:
            raise ValueError("Selected shape is not an editable text target")
        target = (slide_index, shape_id, change.get("row"), change.get("column"))
        # Disallow overlapping edits to make every expected hash refer to source content.
        if target in targets:
            raise ValueError("Only one edit per shape or table cell is supported")
        targets.add(target)
        if "paragraph" in change:
            p = _integer(change["paragraph"], "paragraph", 1, len(frame.paragraphs))
            paragraph = frame.paragraphs[p - 1]
            r = _integer(change["run"], "run", 1, len(paragraph.runs))
            run = paragraph.runs[r - 1]
            if "\n" in text:
                raise ValueError("Run replacement must not contain paragraph breaks")
            previous = run.text
        else:
            if not _plain_frame(frame):
                raise ValueError(
                    "Complex text requires a precise paragraph/run locator; whole replacement would discard formatting or fields"
                )
            if "\n" in text:
                raise ValueError("Whole-target replacement must remain a single paragraph")
            paragraph = frame.paragraphs[0]
            run = paragraph.runs[0] if paragraph.runs else None
            previous = frame.text
        if change.get("expected_text_sha256") != _hash(previous):
            raise ValueError("Stale text: expected_text_sha256 does not match the selected target")
        if run is None:
            run = paragraph.add_run()
        font_size = run.font.size or paragraph.font.size
        run.text = text
        if _plain_frame(frame) and font_size is not None:
            _fit(frame.text, width, height, font_size.pt)
        changed_parts[str(slide.part.partname).lstrip("/")] = slide.part
    # Recheck the source before output creation. Publication is the caller's job.
    if source.read_bytes() != original:
        raise ValueError("Source changed while preparing edits")
    with ZipFile(io.BytesIO(original)) as source_zip, ZipFile(output, "w") as target_zip:
        for member in source_zip.infolist():
            payload = (
                changed_parts[member.filename].blob
                if member.filename in changed_parts
                else source_zip.read(member)
            )
            target_zip.writestr(member, payload)
        target_zip.comment = source_zip.comment
    result = inspect(output)
    result.update(
        status="edited",
        edits_applied=len(edits),
        changed_package_parts=sorted(changed_parts),
        preservation="Unchanged ZIP member payloads copied verbatim; selected slide XML reserialized",
    )
    return result
