"""Bounded XLSX authoring and conservative, copy-only editing.

The caller validates workspace paths and atomically publishes the staged output.
No formula is evaluated here. Inspection deliberately never reads cached results.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
import warnings
import zipfile
from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_SHEETS = 16
MAX_CELLS = 50_000
MAX_ROWS = 10_000
MAX_COLUMNS = 256
MAX_TEXT = 2 * 1024 * 1024
MAX_ZIP_BYTES = 32 * 1024 * 1024
MAX_PREVIEW_CELLS = 1_000
MAX_PREVIEW_TEXT = 64 * 1024
_CELL = re.compile(r"([A-Z]{1,3})([1-9][0-9]{0,4})\Z")
_COLOR = re.compile(r"#[0-9a-fA-F]{6}\Z")
_BAD_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]|[\ud800-\udfff]")
_FORMAT_KEYS = {"bold", "italic", "font_color", "fill", "number_format", "align", "wrap"}
_PARTS = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|"
    r"xl/(?:workbook|styles|sharedStrings)\.xml|xl/theme/theme[0-9]+\.xml|"
    r"xl/worksheets/sheet[0-9]+\.xml|xl/tables/table[0-9]+\.xml|"
    r"xl/_rels/workbook\.xml\.rels|xl/worksheets/_rels/sheet[0-9]+\.xml\.rels)\Z"
)
_UNSUPPORTED_ELEMENTS = {
    "extLst",
    "AlternateContent",
    "oleObjects",
    "controls",
    "legacyDrawing",
    "drawing",
    "customSheetViews",
    "customWorkbookViews",
    "pivotTableParts",
    "externalReferences",
    "dataValidations",
    "conditionalFormatting",
    "sheetProtection",
    "workbookProtection",
    "webPublishItems",
    "smartTags",
    "cellSmartTags",
    "scenarios",
    "richValue",
    "ignoredErrors",
    "phoneticPr",
    "rPh",
    "calcChain",
    "webPublishing",
    "fileSharing",
}
_ROOT_CHILDREN = {
    "workbook": {"fileVersion", "workbookPr", "bookViews", "sheets", "definedNames", "calcPr"},
    "worksheet": {
        "sheetPr",
        "dimension",
        "sheetViews",
        "sheetFormatPr",
        "cols",
        "sheetData",
        "sheetCalcPr",
        "autoFilter",
        "mergeCells",
        "hyperlinks",
        "printOptions",
        "pageMargins",
        "pageSetup",
        "headerFooter",
        "tableParts",
    },
    "table": {"autoFilter", "sortState", "tableColumns", "tableStyleInfo"},
}

GUIDE = {
    "create": "sheets: list of {name, rows, cells?, header_row?, column_widths?, freeze_panes?, tables?, page_setup?}",
    "cell": "JSON scalar is literal; {value: scalar, format?: {...}} or {formula: '=SUM(B2:B3)', format?: {...}}",
    "format": "bold/italic/wrap booleans; font_color/fill #RRGGBB; number_format text; align left/center/right",
    "table": "{range: 'A1:B3', name: 'Revenue', style?: 'Table Style Medium 2'}; first row must be distinct literal text headers",
    "page_setup": "{orientation?: 'portrait'|'landscape', fit_to_pages?: [width, height], print_area?: 'A1:J20'}; 0 means unlimited pages on that axis, 1..10 limits pages. Wrap long notes and inspect rendered pages; scaling alone does not prove legibility.",
    "edit": "{expected_source_sha256: inspect.source_sha256, changes: [{sheet, cell: 'B2', value OR formula, format?}]}",
    "limits": {
        "sheets": MAX_SHEETS,
        "cells_and_bounding_area": MAX_CELLS,
        "rows": MAX_ROWS,
        "columns": MAX_COLUMNS,
        "text_bytes": MAX_TEXT,
        "inspection_preview_cells": MAX_PREVIEW_CELLS,
    },
    "formula_functions": "Function names are preserved; the installed calculation engine determines support. Writing a formula does not prove it calculates.",
    "delivery_workflow": "Create, recalculate, check and visually review working copies outside final deliverables. Repair layout and formulas there, then use fs_copy to copy the reviewed workbook to the requested new final filename and check that exact delivered file. Every operation requires a new output path. Reserve the final name until review is complete; do not delete a draft to replace it. Keep specifications and drafts separate from delivered workbooks.",
    "limitations": [
        "Use the separate recalculate command for real calculation when LibreOffice is available. Inspection alone does not trust caches. No macros or external workbook references.",
        "Formula references use local A1 ranges and English functions; named/structured/array formulas are unsupported.",
        "Edit only simple XLSX workbooks. Unsupported parts, rich text, protected sheets, validation and conditional formatting are rejected.",
        "Edits change existing sheet cells; no row/column insertion, renaming or formula-reference rewriting.",
        "Inspection returns bounded cell text previews, complete counts and an editability report; caches are not returned.",
    ],
}


def example_spec() -> dict:
    return {
        "sheets": [
            {
                "name": "Έσοδα",
                "header_row": True,
                "column_widths": [28, 18],
                "freeze_panes": "A2",
                "rows": [
                    ["Πελάτης", "Ποσό"],
                    ["Αθήνα", 120],
                    ["Θεσσαλονίκη", 80],
                    [
                        "Σύνολο",
                        {
                            "formula": "=SUM(B2:B3)",
                            "format": {"number_format": '#,##0.00 "€"', "bold": True},
                        },
                    ],
                ],
                "tables": [{"range": "A1:B3", "name": "Revenue"}],
            }
        ]
    }


def _object(value, allowed: set[str], where: str) -> dict:
    if not isinstance(value, dict) or any(not isinstance(k, str) for k in value):
        raise ValueError(f"{where} must be an object")
    extra = set(value) - allowed
    if extra:
        raise ValueError(f"Unsupported {where} keys: {', '.join(sorted(extra))}")
    return value


def _text(value, where: str, limit: int = 32767) -> str:
    if not isinstance(value, str) or len(value) > limit or _BAD_TEXT.search(value):
        raise ValueError(f"{where} must be valid text of at most {limit} characters")
    return value


def _position(value: str) -> tuple[int, int]:
    if not isinstance(value, str) or not (match := _CELL.fullmatch(value)):
        raise ValueError("Cell references must use uppercase A1 notation")
    column = 0
    for char in match[1]:
        column = column * 26 + ord(char) - 64
    row = int(match[2])
    if row > MAX_ROWS or column > MAX_COLUMNS:
        raise ValueError(f"Cell exceeds {MAX_ROWS} rows / {MAX_COLUMNS} columns")
    return row, column


def _range(value: str) -> tuple[int, int, int, int]:
    if not isinstance(value, str) or len(value.split(":")) != 2:
        raise ValueError("Table range must use A1:B3 notation")
    first, last = value.split(":")
    r1, c1 = _position(first)
    r2, c2 = _position(last)
    if r2 <= r1 or c2 < c1 or (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_CELLS:
        raise ValueError("Table requires a header, data rows and bounded increasing range")
    return r1, c1, r2, c2


def _print_range(value: str) -> tuple[int, int, int, int]:
    if not isinstance(value, str) or len(value.split(":")) != 2:
        raise ValueError("Print area must use A1:B3 notation")
    first, last = value.split(":")
    r1, c1 = _position(first)
    r2, c2 = _position(last)
    if r2 < r1 or c2 < c1 or (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_CELLS:
        raise ValueError("Print area must be a bounded increasing range")
    return r1, c1, r2, c2


def _is_local_print_area(element, workbook) -> bool:
    # Print areas are workbook-defined names, but are layout metadata, not
    # named formula operands. Admit only a bounded range on its owning sheet.
    if element.get("name") != "_xlnm.Print_Area" or set(element.attrib) != {"name", "localSheetId"}:
        return False
    try:
        index = int(element.get("localSheetId", ""))
        names = [sheet.get("name") for sheet in workbook.findall("{*}sheets/{*}sheet")]
        sheet, reference = (element.text or "").rsplit("!", 1)
        if sheet.startswith("'") and sheet.endswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        if not 0 <= index < len(names) or sheet != names[index]:
            return False
        if not re.fullmatch(r"\$?[A-Z]+\$?[1-9][0-9]*:\$?[A-Z]+\$?[1-9][0-9]*", reference):
            return False
        _print_range(reference.replace("$", ""))
        return True
    except (ValueError, TypeError):
        return False


def _formula(value: str) -> str:
    from openpyxl.formula import Tokenizer

    _text(value, "Formula", 8192)
    if not value.startswith("=") or value == "=":
        raise ValueError("Explicit formula must begin with = and contain an expression")
    # Validate local references structurally. Function support belongs to the
    # calculation engine, not an incomplete hand-maintained function whitelist.
    try:
        tokens = Tokenizer(value).items
    except Exception as exc:
        raise ValueError("Invalid formula syntax") from exc
    for token in tokens:
        if token.type == "ARRAY" or token.subtype == "ROW":
            raise ValueError("Array formulas are unsupported")
        if token.type == "OPERAND" and token.subtype == "RANGE":
            ref = token.value
            if any(char in ref for char in "[]|\\"):
                raise ValueError(
                    "External, named and structured formula references are unsupported"
                )
            if "!" in ref:
                sheet, ref = ref.rsplit("!", 1)
                if sheet.startswith("'") and sheet.endswith("'"):
                    sheet = sheet[1:-1].replace("''", "'")
                _sheet_name(sheet)
            for cell in ref.split(":"):
                _position(cell.replace("$", "").upper())
            if len(ref.split(":")) > 2:
                raise ValueError("Unsupported formula range")
        if token.type == "OPERAND" and token.subtype == "ERROR":
            raise ValueError("Formula contains an Excel error reference")
    return value


def _format(value) -> dict:
    value = _object(value, _FORMAT_KEYS, "cell format")
    for key, item in value.items():
        if key in {"bold", "italic", "wrap"} and not isinstance(item, bool):
            raise ValueError(f"{key} must be boolean")
        if key in {"font_color", "fill"} and (
            not isinstance(item, str) or not _COLOR.fullmatch(item)
        ):
            raise ValueError(f"{key} must be a #RRGGBB color")
        if key == "number_format":
            _text(item, "number_format", 128)
        if key == "align" and (
            not isinstance(item, str) or item not in {"left", "center", "right"}
        ):
            raise ValueError("align must be left, center or right")
    return value


def _cell(value) -> tuple[object, bool, dict]:
    style = {}
    if isinstance(value, dict):
        _object(value, {"value", "formula", "format"}, "cell")
        style = _format(value.get("format", {}))
        if ("value" in value) == ("formula" in value):
            raise ValueError("Cell object requires exactly one of value or formula")
        if "formula" in value:
            return _formula(value["formula"]), True, style
        value = value["value"]
    if isinstance(value, str):
        _text(value, "Cell value")
    elif value is not None and not isinstance(value, (int, float, bool)):
        raise ValueError("Cell value must be a JSON string, finite number, boolean or null")
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            finite = math.isfinite(value)
        except OverflowError:
            finite = False
        if not finite:
            raise ValueError("Cell value must be a finite number")
        if isinstance(value, int) and abs(value) > 999_999_999_999_999:
            raise ValueError("Integers longer than 15 digits must be supplied as literal strings")
    return value, False, style


def _sheet_name(value: str) -> str:
    _text(value, "Sheet name", 31)
    if (
        not value
        or re.search(r"[\[\]:*?/\\]", value)
        or value.startswith("'")
        or value.endswith("'")
    ):
        raise ValueError("Invalid sheet name")
    if value.casefold() == "history":
        raise ValueError("History is a reserved sheet name")
    return value


def _bounded(cells: dict) -> tuple[int, int]:
    max_row = max((r for r, _ in cells), default=0)
    max_col = max((c for _, c in cells), default=0)
    if len(cells) > MAX_CELLS or max_row * max_col > MAX_CELLS:
        raise ValueError(f"Sheet bounding area exceeds {MAX_CELLS} cells")
    return max_row, max_col


def _create_spec(spec: dict) -> list[dict]:
    _object(spec, {"sheets"}, "workbook")
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= MAX_SHEETS:
        raise ValueError(f"sheets must contain between 1 and {MAX_SHEETS} worksheets")
    result, names, table_names = [], set(), set()
    total_area = total_text = 0
    for sheet in sheets:
        _object(
            sheet,
            {
                "name",
                "rows",
                "cells",
                "header_row",
                "column_widths",
                "freeze_panes",
                "tables",
                "page_setup",
            },
            "sheet",
        )
        name = _sheet_name(sheet.get("name"))
        if name.casefold() in names:
            raise ValueError("Sheet names must be unique ignoring case")
        names.add(name.casefold())
        cells = {}
        rows = sheet.get("rows", [])
        if not isinstance(rows, list) or len(rows) > MAX_ROWS:
            raise ValueError(f"rows must be a list of at most {MAX_ROWS} rows")
        for r, row in enumerate(rows, 1):
            if not isinstance(row, list) or len(row) > MAX_COLUMNS:
                raise ValueError(f"Each row must contain at most {MAX_COLUMNS} cells")
            for c, value in enumerate(row, 1):
                cells[r, c] = _cell(value)
        extra = sheet.get("cells", {})
        if not isinstance(extra, dict) or len(extra) > MAX_CELLS:
            raise ValueError("cells must be a bounded A1-keyed object")
        for ref, value in extra.items():
            pos = _position(ref)
            if pos in cells:
                raise ValueError("cells must not overlap rows; specify each cell once")
            cells[pos] = _cell(value)
        area = _bounded(cells)
        total_area += area[0] * area[1]
        total_text += sum(
            len(v.encode("utf-8")) for v, _, _ in cells.values() if isinstance(v, str)
        )
        if total_area > MAX_CELLS or total_text > MAX_TEXT:
            raise ValueError("Workbook exceeds total cell/text limits")
        header = sheet.get("header_row", False)
        if not isinstance(header, bool):
            raise ValueError("header_row must be boolean")
        widths = sheet.get("column_widths", [])
        if (
            not isinstance(widths, list)
            or len(widths) > MAX_COLUMNS
            or any(
                isinstance(w, bool) or not isinstance(w, (int, float)) or not 1 <= w <= 100
                for w in widths
            )
        ):
            raise ValueError("column_widths must be a list of widths from 1 to 100")
        freeze = _position(sheet["freeze_panes"]) if "freeze_panes" in sheet else None
        page_setup = _object(
            sheet.get("page_setup", {}), {"orientation", "fit_to_pages", "print_area"}, "page_setup"
        )
        if "orientation" in page_setup and (
            not isinstance(page_setup["orientation"], str)
            or page_setup["orientation"] not in {"portrait", "landscape"}
        ):
            raise ValueError("page_setup.orientation must be portrait or landscape")
        fit = page_setup.get("fit_to_pages")
        if "fit_to_pages" in page_setup and (
            not isinstance(fit, list)
            or len(fit) != 2
            or any(type(value) is not int or not 0 <= value <= 10 for value in fit)
        ):
            raise ValueError("page_setup.fit_to_pages must be [width, height] with integers 0..10")
        print_bounds = (
            _print_range(page_setup["print_area"]) if "print_area" in page_setup else None
        )
        if print_bounds and (print_bounds[2] > area[0] or print_bounds[3] > area[1]):
            raise ValueError("page_setup.print_area must fit existing cells")
        tables = sheet.get("tables", [])
        if not isinstance(tables, list) or len(tables) > 32:
            raise ValueError("tables must be a list of at most 32 tables")
        normalized_tables = []
        for table in tables:
            _object(table, {"range", "name", "style"}, "table")
            bounds = _range(table.get("range"))
            r1, c1, r2, c2 = bounds
            if r2 > area[0] or c2 > area[1]:
                raise ValueError("Table range must fit existing cells")
            tname = table.get("name")
            if (
                not isinstance(tname, str)
                or not re.fullmatch(r"[A-Za-z_][A-Za-z_0-9]{0,63}", tname)
                or re.fullmatch(r"[A-Za-z]+[0-9]+|[RrCc]", tname)
            ):
                raise ValueError("Table name must be a non-cell ASCII identifier")
            if tname.casefold() in table_names:
                raise ValueError("Table names must be unique ignoring case")
            table_names.add(tname.casefold())
            style = table.get("style", "Table Style Medium 2")
            if not isinstance(style, str) or not re.fullmatch(
                r"Table Style (Light (?:[1-9]|1[0-9]|2[01])|Medium (?:[1-9]|1[0-9]|2[0-8])|Dark (?:[1-9]|1[01]))",
                style,
            ):
                raise ValueError("Unsupported table style")
            headers = [cells.get((r1, c), (None, False, {})) for c in range(c1, c2 + 1)]
            if any(not isinstance(v, str) or not v or f or len(v) > 255 for v, f, _ in headers):
                raise ValueError(
                    "Table headers must be nonempty literal text, at most 255 characters"
                )
            if len({v.casefold() for v, _, _ in headers}) != len(headers):
                raise ValueError("Table headers must be unique ignoring case")
            for other in normalized_tables:
                a, b, c, d = other["bounds"]
                if not (r2 < a or r1 > c or c2 < b or c1 > d):
                    raise ValueError("Table ranges must not overlap")
            normalized_tables.append(
                {
                    "bounds": bounds,
                    "name": tname,
                    "style": style,
                    "headers": [v for v, _, _ in headers],
                }
            )
        # Resolve explicit formula sheet references before creating an output in a later pass.
        result.append(
            {
                "name": name,
                "cells": cells,
                "widths": widths,
                "freeze": freeze,
                "header": header,
                "tables": normalized_tables,
                "page_setup": page_setup,
            }
        )
    for sheet in result:
        for value, formula, _ in sheet["cells"].values():
            if formula:
                _formula_sheets(value, names)
    return result


def _formula_sheets(value: str, names: set[str]) -> None:
    from openpyxl.formula import Tokenizer

    for token in Tokenizer(value).items:
        if token.type == "OPERAND" and token.subtype == "RANGE" and "!" in token.value:
            sheet = token.value.rsplit("!", 1)[0]
            if sheet.startswith("'") and sheet.endswith("'"):
                sheet = sheet[1:-1].replace("''", "'")
            if sheet.casefold() not in names:
                raise ValueError(f"Formula refers to missing sheet: {sheet}")


def create(spec: dict, output: Path, workspace: Path) -> dict:
    import xlsxwriter

    sheets = _create_spec(spec)
    with output.open("xb") as stream:
        with xlsxwriter.Workbook(
            stream,
            {
                "strings_to_formulas": False,
                "strings_to_urls": False,
                "strings_to_numbers": False,
                "in_memory": True,
            },
        ) as book:
            book.set_calc_mode("auto")
            for sheet in sheets:
                ws = book.add_worksheet(sheet["name"])
                page_setup = sheet["page_setup"]
                if page_setup.get("orientation") == "landscape":
                    ws.set_landscape()
                elif page_setup.get("orientation") == "portrait":
                    ws.set_portrait()
                if "fit_to_pages" in page_setup:
                    ws.fit_to_pages(*page_setup["fit_to_pages"])
                if "print_area" in page_setup:
                    ws.print_area(page_setup["print_area"])
                for index, width in enumerate(sheet["widths"]):
                    ws.set_column(index, index, width)
                if sheet["freeze"]:
                    row, col = sheet["freeze"]
                    ws.freeze_panes(row - 1, col - 1)
                for (r, c), (value, formula, style) in sheet["cells"].items():
                    style = dict(style)
                    if sheet["header"] and r == 1:
                        style = {"bold": True, "fill": "#183153", "font_color": "#FFFFFF", **style}
                    mapped = {
                        {
                            "fill": "bg_color",
                            "number_format": "num_format",
                            "wrap": "text_wrap",
                        }.get(k, k): v
                        for k, v in style.items()
                    }
                    fmt = book.add_format(mapped) if mapped else None
                    if formula:
                        # Empty cache avoids presenting XlsxWriter's default zero as a result.
                        ws.write_formula(r - 1, c - 1, value, fmt, "")
                    elif isinstance(value, str):
                        ws.write_string(r - 1, c - 1, value, fmt)
                    elif value is None:
                        ws.write_blank(r - 1, c - 1, None, fmt)
                    else:
                        ws.write(r - 1, c - 1, value, fmt)
                for table in sheet["tables"]:
                    r1, c1, r2, c2 = table["bounds"]
                    ws.add_table(
                        r1 - 1,
                        c1 - 1,
                        r2 - 1,
                        c2 - 1,
                        {
                            "name": table["name"],
                            "style": table["style"],
                            "columns": [{"header": h} for h in table["headers"]],
                        },
                    )
    return {**inspect(output), "status": "created"}


def _snapshot(source: Path) -> bytes:
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx workbooks are supported; macro-enabled files are unsupported")
    if source.stat().st_size > MAX_ZIP_BYTES:
        raise ValueError("Workbook exceeds compressed size limit")
    return source.read_bytes()


def _package(data: bytes) -> list[str]:
    """Check bounds before openpyxl parses the package; report edit-fidelity blockers."""
    blockers = set()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 512 or sum(i.file_size for i in infos) > MAX_ZIP_BYTES:
                raise ValueError("Workbook exceeds package expansion limit")
            if len({i.filename for i in infos}) != len(infos):
                raise ValueError("Workbook contains duplicate package parts")
            sheet_count = cell_count = total_area = total_text = 0
            merged_area = element_count = 0
            for info in infos:
                if info.flag_bits & 1:
                    raise ValueError("Encrypted workbook parts are unsupported")
                name = info.filename
                if not _PARTS.fullmatch(name):
                    blockers.add(f"Unsupported package part: {name}")
                if not name.endswith((".xml", ".rels")):
                    continue
                raw = archive.read(info)
                if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper() or b"\x00" in raw:
                    raise ValueError(
                        "Workbook XML must not contain DTDs, entities or alternate encodings"
                    )
                root = ET.fromstring(raw)
                if (
                    name.startswith("xl/")
                    and not name.endswith(".rels")
                    and not name.startswith("xl/theme/")
                ):
                    if any(
                        not e.tag.startswith(
                            "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                        )
                        for e in root.iter()
                    ):
                        blockers.add("Unsupported spreadsheet XML namespace")
                root_tag = root.tag.rsplit("}", 1)[-1]
                if root_tag in _ROOT_CHILDREN:
                    for child in root:
                        child_tag = child.tag.rsplit("}", 1)[-1]
                        if child_tag not in _ROOT_CHILDREN[root_tag]:
                            blockers.add(f"Unsupported {root_tag} feature: {child_tag}")
                if name == "[Content_Types].xml":
                    if any(
                        "macroenabled" in str(e.attrib).lower() or "vba" in str(e.attrib).lower()
                        for e in root
                    ):
                        blockers.add("Macro-enabled content is unsupported")
                sheet = bool(re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", name))
                positions = {}
                extent_row = extent_col = 0
                if sheet:
                    sheet_count += 1
                for element in root.iter():
                    element_count += 1
                    if element_count > 500_000:
                        raise ValueError("Workbook exceeds XML element limit")
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag in _UNSUPPORTED_ELEMENTS:
                        blockers.add(f"Unsupported workbook feature: {tag}")
                    if (
                        tag == "Relationship"
                        and element.get("TargetMode", "").lower() == "external"
                    ):
                        blockers.add("External relationships are unsupported")
                    if tag == "definedName" and not _is_local_print_area(element, root):
                        blockers.add(
                            "Defined names other than local print areas are unsupported for editing"
                        )
                    if tag == "r" and name == "xl/sharedStrings.xml":
                        blockers.add("Rich cell text is unsupported for editing")
                    if tag in {"t", "f", "v"}:
                        total_text += len((element.text or "").encode("utf-8"))
                    if tag == "f":
                        if element.get("t", "normal") != "normal":
                            blockers.add("Shared, array and data-table formulas are unsupported")
                        if element.text:
                            try:
                                _formula("=" + element.text)
                            except ValueError:
                                blockers.add("Workbook contains unsupported formulas")
                    if sheet and tag == "c":
                        pos = _position(element.get("r", ""))
                        if pos in positions:
                            raise ValueError("Worksheet contains duplicate cells")
                        positions[pos] = None
                        if any(e.tag.rsplit("}", 1)[-1] == "r" for e in element.iter()):
                            blockers.add("Rich cell text is unsupported for editing")
                    if sheet and tag == "dimension":
                        for ref in element.get("ref", "A1").split(":"):
                            row, col = _position(ref)
                            extent_row, extent_col = max(row, extent_row), max(col, extent_col)
                    if tag == "mergeCell":
                        refs = element.get("ref", "").split(":")
                        if len(refs) != 2:
                            raise ValueError("Invalid merged-cell range")
                        r1, c1 = _position(refs[0])
                        r2, c2 = _position(refs[1])
                        if r2 < r1 or c2 < c1:
                            raise ValueError("Invalid merged-cell range")
                        merged_area += (r2 - r1 + 1) * (c2 - c1 + 1)
                        if merged_area > MAX_CELLS:
                            raise ValueError("Workbook exceeds merged-cell area limit")
                        extent_row, extent_col = max(r2, extent_row), max(c2, extent_col)
                    if tag == "calculatedColumnFormula" or tag == "totalsRowFormula":
                        blockers.add("Table formulas are unsupported for editing")
                if sheet:
                    max_row, max_col = _bounded(positions)
                    total_area += max(max_row, extent_row) * max(max_col, extent_col)
                    cell_count += len(positions)
            if (
                not 1 <= sheet_count <= MAX_SHEETS
                or cell_count > MAX_CELLS
                or total_area > MAX_CELLS
                or total_text > MAX_TEXT
            ):
                raise ValueError("Workbook exceeds sheet/cell/text limits")
    except (zipfile.BadZipFile, ET.ParseError, RuntimeError) as exc:
        raise ValueError("Invalid XLSX package") from exc
    return sorted(blockers)


def _load(data: bytes):
    from openpyxl import load_workbook

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        try:
            book = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
        except Exception as exc:
            raise ValueError("Workbook cannot be safely read by openpyxl") from exc
    notes = ["openpyxl reported unsupported or lossy content"] if caught else []
    if any(
        isinstance(cell.value, float) and not math.isfinite(cell.value)
        for ws in book.worksheets
        for row in ws.iter_rows()
        for cell in row
    ):
        notes.append("Non-finite numeric values are unsupported for editing")
    return book, notes


def _json_value(value):
    if isinstance(value, float) and not math.isfinite(value):
        return {"non_finite_number": str(value)}
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {"duration_seconds": value.total_seconds()}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return {"unsupported_type": type(value).__name__}


def _inspection(data: bytes, book, blockers: list[str]) -> dict:
    preview_left, text_left = MAX_PREVIEW_CELLS, MAX_PREVIEW_TEXT
    sheets, total_cells, formula_count = [], 0, 0
    for ws in book.worksheets:
        # Preflight bounded both dimensions and area; iter_rows cannot expand an
        # adversarial A1:XFD1048576 worksheet here.
        cells, count, formulas = [], 0, 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                count += 1
                is_formula = cell.data_type == "f"
                formulas += is_formula
                if preview_left <= 0 or text_left <= 0:
                    continue
                value = _json_value(cell.value)
                truncated = isinstance(value, str) and len(value) > min(512, text_left)
                if isinstance(value, str):
                    value = value[: min(512, text_left)]
                    text_left -= len(value)
                cells.append(
                    {
                        "cell": cell.coordinate,
                        "formula" if is_formula else "value": value,
                        "data_type": cell.data_type,
                        "number_format": cell.number_format,
                        **({"text_truncated": True} if truncated else {}),
                    }
                )
                preview_left -= 1
        sheets.append(
            {
                "name": ws.title,
                "rows": ws.max_row,
                "columns": ws.max_column,
                "cell_count": count,
                "formula_count": formulas,
                "cells": cells,
                "cells_truncated": len(cells) < count,
                "tables": [{"name": t.name, "range": t.ref} for t in ws.tables.values()],
            }
        )
        total_cells += count
        formula_count += formulas
    return {
        "format": "xlsx",
        "source_sha256": hashlib.sha256(data).hexdigest(),
        "sheet_count": len(sheets),
        "cell_count": total_cells,
        "formula_count": formula_count,
        "sheets": sheets,
        "editable": not blockers,
        "edit_blockers": sorted(set(blockers)),
        "needs_recalculation": bool(formula_count),
        "formula_results": "not_calculated",
        "cached_values": "not_read_or_trusted",
        "visual_review": "not_performed",
    }


def inspect(source: Path) -> dict:
    data = _snapshot(source)
    blockers = _package(data)
    book, notes = _load(data)
    try:
        return {**_inspection(data, book, blockers + notes), "status": "inspected"}
    finally:
        book.close()


def _apply_style(cell, style: dict) -> None:
    from openpyxl.styles import PatternFill

    if any(key in style for key in {"bold", "italic", "font_color"}):
        font = copy(cell.font)
        if "bold" in style:
            font.bold = style["bold"]
        if "italic" in style:
            font.italic = style["italic"]
        if "font_color" in style:
            font.color = style["font_color"][1:]
        cell.font = font
    if "fill" in style:
        cell.fill = PatternFill("solid", fgColor=style["fill"][1:])
    if "number_format" in style:
        cell.number_format = style["number_format"]
    if "align" in style or "wrap" in style:
        alignment = copy(cell.alignment)
        if "align" in style:
            alignment.horizontal = style["align"]
        if "wrap" in style:
            alignment.wrap_text = style["wrap"]
        cell.alignment = alignment


def edit(source: Path, spec: dict, output: Path, workspace: Path) -> dict:
    from openpyxl.cell.cell import MergedCell
    from openpyxl.workbook.properties import CalcProperties

    _object(spec, {"expected_source_sha256", "changes"}, "edit")
    expected = spec.get("expected_source_sha256")
    if not isinstance(expected, str) or not re.fullmatch(r"[0-9a-f]{64}", expected):
        raise ValueError("expected_source_sha256 from inspect is required")
    if source.resolve() == output.resolve() or (output.exists() and source.samefile(output)):
        raise ValueError("Editing must write a separate output; the original is never overwritten")
    data = _snapshot(source)
    if hashlib.sha256(data).hexdigest() != expected:
        raise ValueError("Source is stale: expected_source_sha256 does not match")
    blockers = _package(data)
    if blockers:
        raise ValueError("Workbook cannot be safely edited: " + "; ".join(blockers[:8]))
    changes = spec.get("changes")
    if not isinstance(changes, list) or not 1 <= len(changes) <= MAX_CELLS:
        raise ValueError("changes must be a nonempty bounded list")
    book, notes = _load(data)
    try:
        if notes:
            raise ValueError("Workbook cannot be safely edited: " + "; ".join(notes))
        names = {s.casefold() for s in book.sheetnames}
        normalized, touched = [], set()
        for change in changes:
            _object(change, {"sheet", "cell", "value", "formula", "format"}, "change")
            sheet, ref = change.get("sheet"), change.get("cell")
            if not isinstance(sheet, str) or sheet not in book.sheetnames:
                raise ValueError("Change must name an existing worksheet exactly")
            _position(ref)
            if (sheet, ref) in touched:
                raise ValueError("Specify each changed cell once")
            touched.add((sheet, ref))
            value, formula, style = _cell(
                {k: v for k, v in change.items() if k not in {"sheet", "cell"}}
            )
            if formula:
                _formula_sheets(value, names)
            ws = book[sheet]
            cell = ws[ref]
            if isinstance(cell, MergedCell):
                raise ValueError("Cannot change a merged cell except its top-left cell")
            for table in ws.tables.values():
                r1, c1, r2, c2 = _range(table.ref)
                if r1 == cell.row and c1 <= cell.column <= c2:
                    raise ValueError(
                        "Editing table header cells is unsupported; keep table metadata consistent"
                    )
            normalized.append((cell, value, formula, style))
        total_area = sum(ws.max_row * ws.max_column for ws in book.worksheets)
        if total_area > MAX_CELLS:
            raise ValueError("Edited workbook exceeds total cell area limit")
        for cell, value, formula, style in normalized:
            cell.value = value
            if isinstance(value, str) and not formula:
                cell.data_type = "s"
            _apply_style(cell, style)
        total_text = sum(
            len(cell.value.encode("utf-8"))
            for ws in book.worksheets
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
        )
        if total_text > MAX_TEXT:
            raise ValueError("Edited workbook exceeds total text limit")
        book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
        # A source changed during validation is also stale. The loaded bytes remain
        # isolated from the original, and no original file is ever opened for writes.
        if hashlib.sha256(_snapshot(source)).hexdigest() != expected:
            raise ValueError("Source is stale: changed during edit")
        with output.open("xb") as stream:
            book.save(stream)
    finally:
        book.close()
    result = inspect(output)
    return {
        **result,
        "status": "edited",
        "original_sha256": expected,
        "changes_applied": len(changes),
    }
